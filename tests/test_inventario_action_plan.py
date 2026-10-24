"""Offline contract, persistence, UI, export and download regressions."""

import asyncio
import json
from dataclasses import replace
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import MagicMock, patch
from urllib.error import HTTPError, URLError

import flet as ft
from openpyxl import load_workbook

from automations.inventario.action_plan import ActionStore, AuditApi, ManualFields, actions_from_response
from automations.inventario.action_export import export_actions_excel, export_actions_pdf
from automations.inventario.action_panel import main
from tests.test_inventario import walk


def payload():
    item = {"hotel": "Carmel Wind", "data": "30/07/2026", "pontos_de_melhoria": "Organizar uniformes.\nRevisar cadastro.",
            "carimbo_de_data_hora": "17/08/2026 13:16:19", "endereco_de_e_mail": "auditor@example.test"}
    return {"columns": [{"key": key, "label": key} for key in item], "items": [item], "total": 1}


def action():
    return actions_from_response(payload(), "http://localhost/api/auditorias")[0]


class ActionPlanTests(TestCase):
    def test_default_api_initializes_without_optional_environment_variables(self):
        with patch.dict("os.environ", {}, clear=True), patch("automations.inventario.action_plan.load_environment"):
            api = AuditApi()
        self.assertEqual(api.url, "https://api-relatorio-auditoria-inventario-387905886048.europe-west1.run.app/api/auditorias")
        self.assertEqual(api.source_id, "http://127.0.0.1:8001/api/auditorias")

    def test_api_mapping_stable_when_rows_move_and_updates_do_not_change_identity(self):
        original = payload()
        mapped = actions_from_response(original, "source")
        original["values"] = {"999": original["items"][0]}
        original["items"][0]["pontos_de_melhoria"] = "Correção do texto de melhoria"
        updated = actions_from_response(original, "source")
        self.assertEqual(mapped[0].id, updated[0].id)
        self.assertNotEqual(mapped[0].improvement, updated[0].improvement)
        self.assertNotEqual(mapped[0].id, actions_from_response(original, "another source")[0].id)
        original["items"][0]["pontos_de_melhoria"] = " "
        self.assertEqual(actions_from_response(original, "source"), ())

    def test_invalid_or_duplicate_api_response_is_rejected(self):
        for invalid in ({}, {"items": []}, {"columns": [], "items": payload()["items"]}):
            with self.assertRaises(ValueError):
                actions_from_response(invalid, "source")
        data = payload()
        data["items"].append(data["items"][0].copy())
        with self.assertRaisesRegex(ValueError, "duplicada"):
            actions_from_response(data, "source")
        self.assertEqual(actions_from_response({"items": [], "columns": []}, "source"), ())

    def test_http_client_actually_requests_endpoint_and_hides_errors(self):
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps(payload()).encode()
        with patch("automations.inventario.action_plan.urlopen", return_value=response) as request:
            self.assertEqual(len(AuditApi("http://127.0.0.1:8001/api/auditorias").fetch()), 1)
        self.assertEqual(request.call_args.args[0].full_url, "http://127.0.0.1:8001/api/auditorias")
        self.assertEqual(request.call_args.kwargs["timeout"], 30)
        response.__exit__.assert_called_once()
        for error, message in ((HTTPError("url", 503, "secret", {}, None), "HTTP 503"),
                               (URLError("secret"), "Não foi possível")):
            with patch("automations.inventario.action_plan.urlopen", side_effect=error):
                with self.assertRaisesRegex(RuntimeError, message) as raised:
                    AuditApi().fetch()
                self.assertNotIn("secret", str(raised.exception))
        with self.assertRaises(ValueError):
            AuditApi("http://remote.example.test/api").fetch()

    def test_cloud_api_base_url_and_saved_identity_are_compatible_with_local_api(self):
        response = MagicMock()
        response.__enter__.return_value.read.return_value = json.dumps(payload()).encode()
        with patch("automations.inventario.action_plan.urlopen", return_value=response) as request, \
             patch.dict("os.environ", {"INVENTARIO_AUDITORIA_SOURCE_ID": "http://127.0.0.1:8001/api/auditorias"}):
            cloud = AuditApi("https://api.example.test/")
            received = cloud.fetch()
            self.assertEqual(cloud.url, "https://api.example.test/api/auditorias")
            self.assertEqual(received[0].id, actions_from_response(payload(), "http://127.0.0.1:8001/api/auditorias")[0].id)
            self.assertEqual(request.call_args.args[0].full_url, cloud.url)
            explicit = AuditApi("https://api.example.test/api/auditorias")
            self.assertEqual(explicit.fetch()[0].id, received[0].id)

    def test_manual_fields_dates_progress_and_status(self):
        today = date(2026, 10, 22)
        self.assertEqual(ManualFields().status(today), "Não iniciada")
        self.assertEqual(ManualFields(start="01/10/2026").status(today), "Em andamento")
        self.assertEqual(ManualFields(due="20/10/2026").status(today), "Atrasada")
        self.assertEqual(ManualFields(due="20/10/2026", revised="30/10/2026").status(today), "Replanejada")
        self.assertEqual(ManualFields(progress=100).status(today), "Concluída")
        for invalid in (ManualFields(progress=101), ManualFields(start="x"),
                        ManualFields(start="22/10/2026", due="21/10/2026"),
                        ManualFields(completed="22/10/2026", progress=10)):
            with self.assertRaises(ValueError):
                invalid.validate()

    def test_invalid_json_and_oversize_responses_are_rejected(self):
        response = MagicMock()
        response.__enter__.return_value.read.return_value = b"invalid json"
        with patch("automations.inventario.action_plan.urlopen", return_value=response):
            with self.assertRaisesRegex(RuntimeError, "JSON"):
                AuditApi().fetch()
        response.__enter__.return_value.read.return_value = b"123456"
        with patch("automations.inventario.action_plan.urlopen", return_value=response), \
             patch("automations.inventario.action_plan.MAX_RESPONSE_BYTES", 5):
            with self.assertRaisesRegex(ValueError, "8 MB"):
                AuditApi().fetch()

    def test_store_import_is_atomic_and_no_file_is_created_by_panel_construction(self):
        import sqlite3
        with TemporaryDirectory() as folder:
            path = Path(folder) / "store.sqlite3"
            store = ActionStore(path)
            main(MagicMock(web=True), store=store, provider=MagicMock())
            self.assertFalse(path.exists())
            with self.assertRaises(sqlite3.IntegrityError):
                store.sync((action(), replace(action(), id="bad", hotel=None)))
            self.assertEqual(store.list(), ())

    def test_sqlite_survives_restart_sync_keeps_manual_and_conflicts_are_rejected(self):
        with TemporaryDirectory() as folder:
            path = Path(folder) / "actions.sqlite3"
            first = ActionStore(path)
            first.sync((action(),))
            first.save(first.list()[0], ManualFields(how="Organizar por setor", responsible="Responsável", progress=20))
            restarted = ActionStore(path)
            restarted.sync((replace(action(), improvement="Texto atualizado da API"),))
            saved = restarted.list()[0]
            self.assertEqual(saved.manual.how, "Organizar por setor")
            self.assertEqual(saved.improvement, "Texto atualizado da API")
            self.assertEqual(saved.revision, 1)
            self.assertEqual(len(restarted.list()), 1)
            with self.assertRaisesRegex(ValueError, "outra sessão"):
                restarted.save(action(), ManualFields())
            restarted.sync(())
            self.assertEqual(len(restarted.list()), 1)

    def test_excel_types_formulas_and_input_injection(self):
        item = replace(action(), hotel="=1+1", manual=ManualFields(how="=HYPERLINK(\"x\")", responsible="Nome",
            start="01/10/2026", due="30/10/2026", progress=50, notes="@texto"))
        book = load_workbook(BytesIO(export_actions_excel((item,))))
        try:
            sheet = book.active
            self.assertEqual(sheet["A9"].data_type, "s")
            self.assertEqual(sheet["D9"].data_type, "s")
            self.assertEqual(sheet["J9"].value, .5)
            self.assertEqual(sheet["J9"].number_format, "0%")
            self.assertEqual(sheet["F9"].value.date(), date(2026, 10, 1))
            self.assertEqual(sheet["L9"].data_type, "f")
            self.assertIn('"Concluída"', sheet["L9"].value)
            self.assertEqual(sheet.freeze_panes, "C9")
            self.assertEqual(sheet.auto_filter.ref, "A8:M9")
            self.assertEqual(sheet["C9"].value, item.improvement)
            self.assertEqual(sheet["P1"].value, "=TODAY()")
            self.assertTrue(book.calculation.fullCalcOnLoad)
        finally:
            book.close()

    def test_pdf_is_readable_and_long_text_spans_pages_without_loss(self):
        import pypdfium2 as pdfium
        item = replace(action(), improvement=("Melhoria com acentos e detalhes. " * 700) + "MARCADOR FINAL",
                       manual=ManualFields(how="<b>Texto literal</b> & instruções", progress=20))
        content = export_actions_pdf((item,))
        self.assertTrue(content.startswith(b"%PDF-"))
        with pdfium.PdfDocument(content) as document:
            self.assertGreater(len(document), 1)
            pages = []
            for index in range(len(document)):
                page = document[index]
                textpage = page.get_textpage()
                pages.append(textpage.get_text_range())
                textpage.close()
                page.close()
            self.assertIn("MARCADOR FINAL", " ".join(pages))
            self.assertIn("Texto literal", " ".join(pages))

    def test_ui_http_sync_save_reload_filter_and_both_downloads(self):
        with TemporaryDirectory() as folder:
            store = ActionStore(Path(folder) / "actions.sqlite3")
            source = MagicMock()
            source.fetch.return_value = (action(), replace(action(), id="second", hotel="Magna"))
            page, download = MagicMock(web=True), MagicMock(return_value="/downloads/file")
            view = main(page, download_file=download, provider=source, store=store)
            page.update.side_effect = lambda: view.control._build_add_commands()
            source.fetch.assert_not_called()
            refresh = next(c for c in walk(view.control) if isinstance(c, ft.ElevatedButton))
            asyncio.run(refresh.on_click(None))
            grid = next(c for c in walk(view.control) if isinstance(c, ft.DataTable))
            edit = grid.rows[0].cells[-1].content
            edit.on_click(None)
            dialog = page.open.call_args.args[0]
            fields = {c.label: c for c in walk(dialog) if isinstance(c, ft.TextField)}
            fields["Como?"].value, fields["% concluído"].value = "Plano manual", "25"
            asyncio.run(dialog.actions[-1].on_click(None))
            self.assertEqual(store.list()[0].manual.how, "Plano manual")
            asyncio.run(refresh.on_click(None))
            self.assertEqual(store.list()[0].manual.how, "Plano manual")
            company = next(c for c in walk(view.control) if isinstance(c, ft.Dropdown) and c.label == "Hotel")
            company.value = "Carmel Wind"
            company.on_change(None)
            view.apply_theme(False)
            for label, suffix in (("Baixar Excel", ".xlsx"), ("Baixar PDF", ".pdf")):
                button = next(c for c in walk(view.control) if isinstance(c, ft.OutlinedButton) and c.text == label)
                asyncio.run(button.on_click(None))
                self.assertTrue(download.call_args.args[1].endswith(suffix))
            self.assertEqual(len(next(c for c in walk(view.control) if isinstance(c, ft.DataTable)).rows), 1)
            source.fetch.side_effect = RuntimeError("API indisponível")
            with patch("automations.inventario.action_panel.logging.getLogger"):
                asyncio.run(refresh.on_click(None))
            self.assertTrue(any(isinstance(c, ft.DataTable) for c in walk(view.control)))
            second = main(page, provider=source, store=ActionStore(store.path))
            load = next(c for c in walk(second.control) if isinstance(c, ft.OutlinedButton) and c.text == "Carregar salvos")
            asyncio.run(load.on_click(None))
            self.assertEqual(len(next(c for c in walk(second.control) if isinstance(c, ft.DataTable)).rows), 2)

    def test_export_downloads_are_single_use_and_route_is_served(self):
        from fastapi.testclient import TestClient
        from portal.web import Downloads, create_app
        store = Downloads()
        path = store.issue_file(export_actions_pdf((action(),)), "plano.pdf")
        with patch("portal.web.Downloads", return_value=store), TestClient(create_app()) as client:
            self.assertEqual(client.get("/inventario/plano-acao").status_code, 200)
            response = client.get(path)
            self.assertTrue(response.content.startswith(b"%PDF-"))
            self.assertEqual(response.headers["cache-control"], "no-store")
            self.assertEqual(client.get(path).status_code, 404)

    def test_portal_action_route_and_theme_preserve_loaded_data(self):
        from portal.app import main as portal
        with TemporaryDirectory() as folder:
            store = ActionStore(Path(folder) / "store.sqlite3")
            page = MagicMock(web=True, route="/inventario/plano-acao")
            roots = []
            page.add.side_effect = lambda *items: roots.extend(items)
            with patch("automations.inventario.action_panel.ActionStore", return_value=store), \
                 patch("automations.inventario.action_panel.AuditApi") as source:
                source.return_value.fetch.return_value = (action(),)
                portal(page, MagicMock(), MagicMock())
                sidebar, body = roots[0].controls
                original = body.content
                refresh = next(c for c in walk(body) if isinstance(c, ft.ElevatedButton))
                asyncio.run(refresh.on_click(None))
                toggle = next(c for c in walk(sidebar) if isinstance(c, ft.TextButton) and c.text == "Modo claro")
                toggle.on_click(None)
                page.route = "/bordero"
                page.on_route_change(None)
                page.route = "/inventario/plano-acao"
                page.on_route_change(None)
                self.assertIs(body.content, original)
                self.assertTrue(any(isinstance(c, ft.DataTable) for c in walk(body)))
                source.return_value.fetch.assert_called_once()
