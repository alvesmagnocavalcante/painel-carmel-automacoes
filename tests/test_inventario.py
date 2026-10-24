"""Offline inventory, SQL, export and Flet serialization regressions."""

import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest import TestCase
from unittest.mock import MagicMock, patch

import flet as ft
import psycopg
from openpyxl import load_workbook

from automations.inventario.data import InventoryCount, build_matrix, months_between
from automations.inventario.database import SqlInventoryProvider, connection_settings
from automations.inventario.export import export_inventory
from automations.inventario.panel import main
from automations.inventario.charts import company_chart
from portal.theme import DARK, LIGHT


JAN = date(2026, 1, 1)
FEB = date(2026, 2, 1)


def counts():
    return (InventoryCount(JAN, "TAÍBA", "BAR", 100, 10),
            InventoryCount(JAN, "TAÍBA", "COZINHA", 10, 10),
            InventoryCount(FEB, "CHARME", "BAR", 20, 0))


def walk(control):
    yield control
    content = getattr(control, "content", None)
    if isinstance(content, ft.Control):
        yield from walk(content)
    for child in getattr(control, "controls", []) or []:
        yield from walk(child)


def month_fields(view):
    next(item for item in walk(view.control) if isinstance(item, ft.Checkbox)).value = True
    return [item for item in walk(view.control) if isinstance(item, ft.TextField)]


class InventoryTests(TestCase):
    def test_chart_uses_weighted_percentages_and_theme_without_io(self):
        for colors in (DARK, LIGHT):
            view = company_chart(counts(), colors)
            chart = next(item for item in walk(view) if isinstance(item, ft.BarChart))
            self.assertEqual(len(chart.bar_groups), 2)
            self.assertEqual(len(chart.top_axis.labels), 2)
            top_values = chart.top_axis.labels[0].label.controls
            self.assertEqual(top_values[0].value, "18,18%")
            self.assertEqual(top_values[1].value, "20 com diferença\n110 contados")
            self.assertEqual(top_values[0].color, colors["warning"])
            self.assertEqual(chart.top_axis.labels[0].value, chart.bar_groups[0].x)
            rod = chart.bar_groups[0].bar_rods[0]
            self.assertAlmostEqual(rod.to_y, 20 / 110 * 100)
            self.assertIn("20 de 110 contados", rod.tooltip)
            self.assertEqual(rod.color, colors["warning"])
            self.assertEqual(chart.min_y, 0)
            self.assertGreaterEqual(chart.max_y, rod.to_y)
            self.assertTrue(view._build_add_commands())
        empty = company_chart((InventoryCount(JAN, "A", "B", 0, 0),), DARK)
        self.assertFalse(any(isinstance(item, ft.BarChart) for item in walk(empty)))

    def test_single_company_chart_ranks_warehouses_with_weighted_percentages(self):
        rows = (counts()[0], counts()[1], InventoryCount(FEB, "TAÍBA", "BAR", 10, 2),
                InventoryCount(FEB, "TAÍBA", "SEM DIFERENÇA", 20, 0),
                InventoryCount(FEB, "TAÍBA", "SEM CONTAGEM", 0, 0))
        for colors in (DARK, LIGHT):
            view = company_chart(rows, colors)
            bars = [item for item in walk(view) if isinstance(item, ft.ProgressBar)]
            self.assertEqual(len(bars), 3)
            self.assertEqual(bars[0].value, 1)
            self.assertIn("COZINHA", bars[0].tooltip)
            self.assertAlmostEqual(bars[1].value, 12 / 110)
            self.assertIn("12 de 110 contados", bars[1].tooltip)
            self.assertEqual(bars[2].value, 0)
            labels = [item.value for item in walk(view) if isinstance(item, ft.Text)]
            self.assertIn("12 com diferença de 110 contados", labels)
            self.assertIn("0 com diferença de 20 contados", labels)
            self.assertEqual(bars[0].color, colors["warning"])
            self.assertEqual(bars[0].bgcolor, colors["warning_soft"])
            self.assertEqual(view.scroll, ft.ScrollMode.AUTO)
            self.assertFalse(any(isinstance(item, ft.BarChart) for item in walk(view)))
            self.assertTrue(view._build_add_commands())

    def test_table_height_stays_fixed_before_and_after_company_filter(self):
        page, provider = MagicMock(web=True, width=1600), MagicMock()
        provider.fetch.return_value = counts()
        view = main(page, provider=provider)
        self.assertTrue(any(isinstance(item, ft.Container) and item.height == 420 for item in walk(view.control)))
        fields = month_fields(view)
        fields[0].value, fields[1].value = "01/2026", "02/2026"
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        asyncio.run(consult.on_click(None))
        company = next(item for item in walk(view.control) if isinstance(item, ft.Dropdown))
        for value in ("Todos", "CHARME", "TAÍBA"):
            company.value = value
            company.on_change(None)
            vertical = next(item for item in walk(view.control) if isinstance(item, ft.Column) and item.controls and isinstance(item.controls[0], ft.DataTable))
            self.assertEqual(vertical.height, 420)
            self.assertEqual(vertical.controls[0].data_row_min_height, 52)
            self.assertEqual(vertical.controls[0].data_row_max_height, 52)
            labels = [item.value for item in walk(view.control) if isinstance(item, ft.Text)]
            self.assertIn("Diferenças por empresa" if value == "Todos" else "Diferenças por almoxarifado", labels)
            self.assertTrue(view.control._build_add_commands())
        provider.fetch.assert_called_once()

    def test_results_fill_width_and_keep_scroll_for_wide_periods(self):
        page, provider = MagicMock(web=True, width=1600), MagicMock()
        provider.fetch.return_value = (counts()[0],)
        view = main(page, provider=provider)
        self.assertEqual(view.control.horizontal_alignment, ft.CrossAxisAlignment.STRETCH)
        fields = month_fields(view)
        fields[0].value, fields[1].value = "01/2026", "01/2026"
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        asyncio.run(consult.on_click(None))
        vertical = next(item for item in walk(view.control) if isinstance(item, ft.Column) and item.controls and isinstance(item.controls[0], ft.DataTable))
        self.assertEqual(vertical.horizontal_alignment, ft.CrossAxisAlignment.STRETCH)
        self.assertEqual(vertical.scroll, ft.ScrollMode.ALWAYS)
        self.assertEqual(vertical.width, 1510)
        self.assertEqual(vertical.controls[0].width, 1510)
        self.assertTrue(view.control._build_add_commands())
        fields[1].value = "02/2026"
        asyncio.run(consult.on_click(None))
        vertical = next(item for item in walk(view.control) if isinstance(item, ft.Column) and item.controls and isinstance(item.controls[0], ft.DataTable))
        self.assertEqual(vertical.width, 1510)
        self.assertEqual(vertical.controls[0].width, 1510)
        self.assertTrue(any(isinstance(item, ft.Row) and item.scroll == ft.ScrollMode.ALWAYS for item in walk(view.control)))
        self.assertTrue(view.control._build_add_commands())

    def test_resize_preserves_table_and_snapshot_without_fetching(self):
        page, provider, width = MagicMock(web=True), MagicMock(), MagicMock(return_value=1200)
        provider.fetch.return_value = counts()
        view = main(page, provider=provider, available_width=width)
        fields = month_fields(view)
        fields[0].value, fields[1].value = "01/2026", "02/2026"
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        asyncio.run(consult.on_click(None))
        table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
        self.assertEqual(table.width, 1200)
        width.return_value = 1500
        self.assertTrue(view.resize())
        self.assertEqual(table.width, 1500)
        self.assertFalse(view.resize())
        width.return_value = 600
        self.assertTrue(view.resize())
        self.assertEqual(table.width, 1040)
        self.assertIs(next(item for item in walk(view.control) if isinstance(item, ft.DataTable)), table)
        provider.fetch.assert_called_once()

    def test_optional_period_defaults_to_all_and_can_be_cleared(self):
        page, provider = MagicMock(web=True), MagicMock()
        provider.fetch.return_value = counts()
        download = MagicMock(return_value="/downloads/inventory")
        view = main(page, provider=provider, download_file=download)
        checkbox = next(item for item in walk(view.control) if isinstance(item, ft.Checkbox))
        self.assertFalse(checkbox.value)
        self.assertTrue(any(isinstance(item, ft.ResponsiveRow) and not item.visible for item in walk(view.control)))
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        asyncio.run(consult.on_click(None))
        provider.fetch.assert_called_once_with(None, None, "Todos")
        self.assertTrue(any(isinstance(item, ft.Text) and "Todos os meses disponíveis •" in str(item.value) for item in walk(view.control)))
        export = next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton))
        export.on_click(None)
        self.assertEqual(download.call_args.args[1], "inventario-20260101-20260228.xlsx")
        fields = month_fields(view)
        checkbox.on_change(None)
        fields[0].value, fields[1].value = "02/2024", "02/2024"
        asyncio.run(consult.on_click(None))
        provider.fetch.assert_called_with(date(2024, 2, 1), date(2024, 2, 29), "Todos")
        clear = next(item for item in walk(view.control) if isinstance(item, ft.TextButton) and item.text == "Limpar período")
        asyncio.run(clear.on_click(None))
        self.assertFalse(checkbox.value)
        provider.fetch.assert_called_with(None, None, "Todos")
        self.assertTrue(view.control._build_add_commands())

    def test_all_history_over_24_months_is_not_truncated(self):
        rows = (InventoryCount(date(2023, 1, 1), "A", "B", 10, 1), counts()[0])
        provider = MagicMock()
        provider.fetch.return_value = rows
        view = main(MagicMock(web=True), provider=provider)
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        asyncio.run(consult.on_click(None))
        table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
        self.assertEqual(len(table.columns), 1 + 37 * 3)
        workbook = load_workbook(BytesIO(export_inventory(rows, date(2023, 1, 1), date(2026, 1, 31), "Todos")))
        try:
            self.assertEqual(workbook.active.max_column, 112)
            self.assertEqual(workbook.active["B6"].value, "jan/23")
        finally:
            workbook.close()

    def test_unfiltered_sql_omits_dates_and_partial_bounds_are_rejected(self):
        connection = MagicMock()
        cursor = connection.__enter__.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = []
        with patch("automations.inventario.database.connection_settings", return_value={}), patch("automations.inventario.database.psycopg.connect", return_value=connection):
            self.assertEqual(SqlInventoryProvider().fetch(), ())
            query, params = cursor.execute.call_args.args
            self.assertNotIn("dtdata >= %s", query)
            self.assertEqual(params, ())
            with self.assertRaises(ValueError):
                SqlInventoryProvider().fetch(JAN, None)
        connection.__exit__.assert_called_once()

    def test_matrix_totals_are_weighted_and_missing_months_are_not_zero(self):
        matrix = build_matrix(counts(), (JAN, FEB))
        taiba = next(line for line in matrix if line.kind == "company" and line.label == "TAÍBA")
        self.assertEqual((taiba.values[0].counted, taiba.values[0].different), (110, 20))
        self.assertEqual(taiba.values[0].percentage, Decimal(20) / Decimal(110))
        self.assertIsNone(taiba.values[1])
        self.assertEqual(matrix[-1].label, "Total geral")
        self.assertEqual(matrix[-1].values[1].different, 0)
        self.assertEqual(matrix[-1].values[1].percentage, Decimal(0))

    def test_invalid_counts_duplicates_and_periods_are_rejected(self):
        with self.assertRaises(ValueError):
            InventoryCount(JAN, "A", "B", 10, 11)
        with self.assertRaises(ValueError):
            build_matrix((counts()[0], counts()[0]), (JAN,))
        with self.assertRaises(ValueError):
            months_between(FEB, JAN)
        self.assertEqual(len(months_between(JAN, date(2028, 1, 1))), 25)
        self.assertEqual(months_between(date(2025, 12, 31), FEB), (date(2025, 12, 1), JAN, FEB))

    def test_query_is_monthly_distinct_parameterized_read_only_and_closes(self):
        connection = MagicMock()
        session = connection.__enter__.return_value
        cursor = session.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [dict(mes=JAN, empresa="TAÍBA", almox="BAR", itens_contados=10, itens_com_diferenca=2)]
        company = "TAÍBA'; DROP TABLE x; --"
        with patch("automations.inventario.database.connection_settings", return_value={}), patch("automations.inventario.database.psycopg.connect", return_value=connection):
            result = SqlInventoryProvider().fetch(JAN, date(2026, 1, 31), company)
        query, params = cursor.execute.call_args.args
        self.assertEqual(params, (JAN, FEB, company))
        self.assertNotIn(company, query)
        self.assertIn("dtdata < %s", query)
        self.assertIn("GROUP BY date_trunc('month', dtdata)::date, empresa, almox", query)
        self.assertIn("COUNT(DISTINCT fkartigo) FILTER (WHERE dif <> 0)", query)
        self.assertTrue(session.read_only)
        cursor.execute.assert_any_call("SET LOCAL statement_timeout = '30s'")
        connection.__exit__.assert_called_once()
        session.cursor.return_value.__exit__.assert_called_once()
        self.assertEqual(result[0].different, 2)

    def test_sql_failure_closes_connection_and_hides_driver_details(self):
        connection = MagicMock()
        cursor = connection.__enter__.return_value.cursor.return_value.__enter__.return_value
        cursor.execute.side_effect = psycopg.OperationalError("private driver details")
        with patch("automations.inventario.database.connection_settings", return_value={}), patch("automations.inventario.database.psycopg.connect", return_value=connection):
            with self.assertRaisesRegex(RuntimeError, "Não foi possível") as raised:
                SqlInventoryProvider().fetch(JAN, FEB)
        self.assertNotIn("private", str(raised.exception))
        connection.__exit__.assert_called_once()

    def test_connection_uses_environment_without_hardcoded_credentials(self):
        variables = {"DB_HOST": "test-host", "DB_NAME": "test-db", "DB_USER": "test-user", "DB_PASS": "test-only",
                     "INVENTARIO_DB_NAME": "inventory-db", "INVENTARIO_ENV_FILE": "test.env"}
        with patch.dict("os.environ", variables, clear=True), patch("automations.inventario.database.load_dotenv") as load:
            settings = connection_settings()
        load.assert_called_once_with("test.env")
        self.assertEqual(settings["dbname"], "inventory-db")
        self.assertEqual(settings["application_name"], "inventario-panel")
        self.assertEqual(settings["connect_timeout"], 10)

    def test_invalid_database_result_is_rejected_after_connection_is_closed(self):
        connection = MagicMock()
        session = connection.__enter__.return_value
        cursor = session.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [dict(mes=JAN, empresa="A", almox="B",
                                             itens_contados=1, itens_com_diferenca=2)]
        with patch("automations.inventario.database.connection_settings", return_value={}), \
             patch("automations.inventario.database.psycopg.connect", return_value=connection):
            with self.assertRaises(ValueError):
                SqlInventoryProvider().fetch()
        connection.__exit__.assert_called_once()
        session.cursor.return_value.__exit__.assert_called_once()

    def test_excel_has_numeric_percentages_same_totals_and_no_formulas(self):
        rows = (*counts(), InventoryCount(JAN, "=1+1", "@warehouse", 1, 0))
        workbook = load_workbook(BytesIO(export_inventory(rows, JAN, date(2026, 2, 28), "Todos")))
        try:
            sheet = workbook.active
            self.assertEqual(sheet["B6"].value, "jan/26")
            self.assertEqual(sheet["E6"].value, "fev/26")
            self.assertEqual(sheet.freeze_panes, "B8")
            totals = sheet.max_row
            self.assertEqual(sheet.cell(totals, 1).value, "Total geral")
            self.assertEqual(sheet.cell(totals, 2).value, 111)
            self.assertEqual(sheet.cell(totals, 3).value, 20)
            self.assertAlmostEqual(sheet.cell(totals, 4).value, 20 / 111)
            self.assertEqual(sheet.cell(totals, 4).number_format, "0.00%")
            self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
            self.assertIn("B6:D6", {str(item) for item in sheet.merged_cells.ranges})
        finally:
            workbook.close()

    def test_panel_consult_filter_export_theme_and_error_preserve_snapshot(self):
        page, provider, download = MagicMock(web=True), MagicMock(), MagicMock(return_value="/downloads/inventory")
        provider.fetch.return_value = counts()
        view = main(page, download_file=download, provider=provider)
        page.update.side_effect = lambda: view.control._build_add_commands()
        provider.fetch.assert_not_called()
        month_fields(view)
        start = next(item for item in walk(view.control) if isinstance(item, ft.TextField) and item.label == "Mês inicial")
        end = next(item for item in walk(view.control) if isinstance(item, ft.TextField) and item.label == "Mês final")
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        start.value, end.value = "01/2026", "02/2026"
        asyncio.run(consult.on_click(None))
        provider.fetch.assert_called_once_with(JAN, date(2026, 2, 28), "Todos")
        company = next(item for item in walk(view.control) if isinstance(item, ft.Dropdown))
        company.value = "TAÍBA"
        company.on_change(None)
        table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
        self.assertEqual(len(table.rows), 4)
        self.assertEqual(table.rows[-1].cells[1].content.value, "110")
        view.apply_theme(False)
        table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
        self.assertEqual(table.rows[-1].cells[0].content.color, "#151311")
        provider.fetch.assert_called_once()
        # Export uses the loaded period, not unsaved edits in the input fields.
        start.value = "01/2027"
        export = next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton))
        export.on_click(None)
        content, name = download.call_args.args
        self.assertEqual(name, "inventario-20260101-20260228.xlsx")
        workbook = load_workbook(BytesIO(content))
        try:
            self.assertEqual(workbook.active["A3"].value, "Empresa: TAÍBA")
        finally:
            workbook.close()
        page.launch_url.assert_called_once_with("/downloads/inventory", web_window_name="_self")
        start.value = "01/2026"
        provider.fetch.side_effect = RuntimeError("Falha controlada")
        with patch("automations.inventario.panel.logging.getLogger"):
            asyncio.run(consult.on_click(None))
        table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
        self.assertEqual(table.rows[-1].cells[1].content.value, "110")
        self.assertFalse(consult.disabled)
        self.assertTrue(view.control._build_add_commands())

    def test_empty_invalid_period_and_desktop(self):
        page, provider = MagicMock(web=True), MagicMock()
        provider.fetch.return_value = ()
        view = main(page, provider=provider)
        page.update.side_effect = lambda: view.control._build_add_commands()
        fields = month_fields(view)
        fields[0].value, fields[1].value = "invalid", "01/2026"
        consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton))
        with patch("automations.inventario.panel.logging.getLogger"):
            asyncio.run(consult.on_click(None))
        provider.fetch.assert_not_called()
        fields[0].value = "01/2026"
        asyncio.run(consult.on_click(None))
        self.assertFalse(any(isinstance(item, ft.DataTable) for item in walk(view.control)))
        self.assertTrue(next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton)).disabled)
        with self.assertRaises(RuntimeError):
            main(MagicMock(web=False))

    def test_inventory_route_and_excel_download_are_served_single_use(self):
        from fastapi.testclient import TestClient
        from portal.web import Downloads, create_app

        store = Downloads()
        payload = export_inventory(counts(), JAN, date(2026, 2, 28), "Todos")
        path = store.issue_file(payload, "inventario.xlsx")
        with patch("portal.web.Downloads", return_value=store), TestClient(create_app()) as client:
            self.assertEqual(client.get("/inventario").status_code, 200)
            response = client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.content, payload)
            self.assertIn("attachment", response.headers["content-disposition"])
            self.assertEqual(response.headers["cache-control"], "no-store")
            self.assertEqual(client.get(path).status_code, 404)
