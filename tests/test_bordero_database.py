"""Offline regression tests for the Bordero source and local reviews."""

from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from dataclasses import replace
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import MagicMock, patch
import asyncio

import flet as ft
from openpyxl import load_workbook

from automations.bordero.database import SqlConferenceProvider, map_document
from automations.bordero.panel_data import classify_conference, export_conferences, filter_conferences, format_tax_id


def document(**changes):
    record = dict(empresa="MAGNA", annumero="123", anhistorico="Serviço",
                  formadepagamento="Transferência", saldo=Decimal("100.50"),
                  dtprogramada=None, dtvencimento=date(2026, 10, 5),
                  usuariolancamento="Ana", numerodeanexos=0, anexos=None,
                  favorecido="Fornecedor exemplo", iddocumentocapcarbase=123,
                  situacaobpm="Aprovado")
    return record | changes


class BorderoTests(TestCase):
    def test_export_preserves_external_text_without_excel_formulas(self):
        fields = ("supplier", "tax_id", "launch_number", "company", "payment_method",
                  "history", "situation", "responsible", "pending_reason", "source_situation")
        values = ('=HYPERLINK("https://example.com","Abrir")', "+123", "-123", "@SUM(A1)", "#N/A")
        rows = tuple(replace(map_document(document()), **dict.fromkeys(fields, value))
                     for value in values)
        with TemporaryDirectory() as directory:
            path = export_conferences(rows, Path(directory) / "literal-text.xlsx",
                                      hotel="=1+1", start_date=date(2026, 10, 1),
                                      end_date=date(2026, 10, 5))
            workbook = load_workbook(path)
            try:
                sheet = workbook.active
                self.assertEqual(sheet["B2"].value, "=1+1")
                self.assertEqual(sheet["B2"].data_type, "s")
                for row_number, value in enumerate(values, 6):
                    for column in (1, 2, 3, 4, 7, 8, 9, 10, 12, 14):
                        with self.subTest(value=value, column=column):
                            cell = sheet.cell(row_number, column)
                            self.assertEqual(cell.value, value)
                            self.assertEqual(cell.data_type, "s")
                    self.assertEqual(sheet.cell(row_number, 5).value, 100.50)
                    self.assertEqual(sheet.cell(row_number, 5).data_type, "n")
                    self.assertEqual(sheet.cell(row_number, 11).value.date(), date(2026, 10, 5))
                    self.assertEqual(sheet.cell(row_number, 11).data_type, "d")
            finally:
                workbook.close()

    def test_search_accepts_supplier_tax_id_and_launch_and_combines_filters(self):
        first = map_document(document(favorecido="Águas do Nordeste", cpfcnpj="00123456000199", annumero="LAN-123"))
        second = replace(map_document(document(favorecido="José Silva", cpfcnpj="01234567890", annumero="456")), situation="Pendente")
        rows = (first, second)
        for query in ("aguas", "  ÁGUAS DO  ", "00123456000199", "00.123.456/0001-99", "000199", "lan-123"):
            with self.subTest(query=query):
                self.assertEqual(filter_conferences(rows, search=query), (first,))
        for query in ("jose", "01234567890", "012.345.678-90", "456"):
            self.assertIn(second, filter_conferences(rows, search=query))
        self.assertEqual(filter_conferences(rows, search="123456"), rows)
        self.assertEqual(filter_conferences(rows, "Pendente", second.payment_method, "jose"), (second,))
        self.assertEqual(filter_conferences(rows, "Conferido", search="jose"), ())
        self.assertEqual(filter_conferences(rows, search="   "), rows)
        self.assertEqual(filter_conferences(rows, search="xyz123"), ())
        self.assertEqual(filter_conferences(rows, search="/"), (first,))

    def test_panel_search_updates_results_export_and_can_be_cleared_without_query(self):
        from automations.bordero.panel import main

        def walk(control):
            yield control
            content = getattr(control, "content", None)
            if isinstance(content, ft.Control):
                yield from walk(content)
            for child in getattr(control, "controls", []) or []:
                yield from walk(child)

        page, download = MagicMock(web=True), MagicMock(return_value="/downloads/test")
        rows = (map_document(document(favorecido="Águas", cpfcnpj="00123456000199")),
                map_document(document(favorecido="José", annumero="456")))
        with patch("automations.bordero.panel.SqlConferenceProvider") as source:
            source.return_value.fetch.return_value = rows
            view = main(page, web_download=download)
            consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton) and item.text.startswith("Consultar"))
            asyncio.run(consult.on_click(None))
            field = next(item for item in walk(view.control) if isinstance(item, ft.TextField) and item.label.startswith("Pesquisar"))
            field.value = "aguas"
            field.on_change(None)
            table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table.rows), 1)
            export = next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton) and item.text == "Baixar Excel")
            export.on_click(None)
            self.assertEqual(download.call_args.args[0], (rows[0],))
            view.toggle_theme(None)
            self.assertEqual(field.color, "#151311")
            self.assertTrue(view.control._build_add_commands())
            field.value = "inexistente"
            field.on_change(None)
            self.assertTrue(export.disabled)
            self.assertTrue(view.control._build_add_commands())
            field.suffix.on_click(None)
            self.assertEqual(field.value, "")
            table = next(item for item in walk(view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table.rows), 2)
            source.return_value.fetch.assert_called_once()

    def test_theme_text_and_semantic_colors_have_sufficient_contrast(self):
        from automations.bordero.panel import LIGHT, DARK

        def luminance(color):
            channels = [int(color[index:index + 2], 16) / 255 for index in (1, 3, 5)]
            linear = [value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4 for value in channels]
            return sum(value * weight for value, weight in zip(linear, (0.2126, 0.7152, 0.0722)))

        for palette in (LIGHT, DARK):
            for key in ("text", "muted", "success", "warning", "info", "money"):
                for background in ("surface", "background"):
                    values = sorted((luminance(palette[key]), luminance(palette[background])))
                    with self.subTest(theme=palette["background"], color=key, background=background):
                        self.assertGreaterEqual((values[1] + 0.05) / (values[0] + 0.05), 4.5)

    def test_cpf_cnpj_formatting_preserves_unknown_values(self):
        cases = (
            ("01234567890", "012.345.678-90"),
            ("012.345.678-90", "012.345.678-90"),
            ("00123456000199", "00.123.456/0001-99"),
            ("00.123.456/0001-99", "00.123.456/0001-99"),
            (None, "Não informado"), ("  ", "Não informado"),
            ("123", "123"), ("ABC01234567890", "ABC01234567890"),
        )
        for value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(format_tax_id(value), expected)
                self.assertEqual(map_document(document(cpfcnpj=value)).tax_id, expected)

    def test_export_contains_cpf_and_cnpj_as_text(self):
        rows = tuple(map_document(document(cpfcnpj=value)) for value in ("01234567890", "00123456000199"))
        with TemporaryDirectory() as directory:
            path = export_conferences(rows, Path(directory) / "documents.xlsx", hotel="Todos", start_date=date(2026, 10, 1), end_date=date(2026, 10, 5))
            workbook = load_workbook(path)
            try:
                self.assertEqual(workbook.active["B5"].value, "CPF/CNPJ")
                for index, row in enumerate(rows, start=6):
                    cell = workbook.active.cell(index, 2)
                    self.assertEqual(cell.value, row.tax_id)
                    self.assertEqual(cell.data_type, "s")
            finally:
                workbook.close()

    def test_export_contains_review_reason_and_due_date(self):
        row = replace(map_document(document()), situation="Pendente", pending_reason="Falta anexo")
        with TemporaryDirectory() as directory:
            output = export_conferences((row,), Path(directory) / "review.xlsx", hotel="Todos", start_date=date(2026, 10, 1), end_date=date(2026, 10, 5))
            workbook = load_workbook(output)
            try:
                sheet = workbook.active
                self.assertEqual(sheet["I6"].value, "Pendente")
                self.assertEqual(sheet["L6"].value, "Falta anexo")
                self.assertEqual(sheet["K6"].value.date(), date(2026, 10, 5))
                self.assertEqual(sheet["M6"].value, "Sem anexo")
                self.assertEqual(sheet["N6"].value, "Aprovado")
            finally:
                workbook.close()

    def test_table_scrollbars_are_bounded_with_many_payments(self):
        from automations.bordero.panel import main

        page = MagicMock(web=True)

        def walk(control):
            yield control
            content = getattr(control, "content", None)
            if isinstance(content, ft.Control):
                yield from walk(content)
            for child in getattr(control, "controls", []) or []:
                yield from walk(child)

        with patch("automations.bordero.panel.SqlConferenceProvider") as source:
            source.return_value.fetch.return_value = tuple(
                map_document(document(iddocumentocapcarbase=index)) for index in range(30)
            )
            view = main(page, web_download=MagicMock())
            consult = next(item for item in walk(view.control) if isinstance(item, ft.ElevatedButton) and item.text.startswith("Consultar"))
            asyncio.run(consult.on_click(None))
            horizontal = next(item for item in walk(view.control) if isinstance(item, ft.Row) and item.scroll == ft.ScrollMode.ALWAYS)
            self.assertEqual(horizontal.height, 420)
            vertical = horizontal.controls[0]
            self.assertIsInstance(vertical, ft.Column)
            self.assertEqual(vertical.height, horizontal.height)
            self.assertEqual(vertical.scroll, ft.ScrollMode.ALWAYS)
            self.assertIsInstance(vertical.controls[0], ft.DataTable)
            self.assertEqual(len(vertical.controls[0].rows), 30)
            source.return_value.fetch.assert_called_once()

    def test_panel_consult_and_review(self):
        from automations.bordero.panel import main

        page = MagicMock()
        page.web = True
        page.overlay = []
        controls = []
        page.add.side_effect = lambda *items: controls.extend(items)

        def walk(control):
            yield control
            content = getattr(control, "content", None)
            if isinstance(content, ft.Control):
                yield from walk(content)
            for child in getattr(control, "controls", []) or []:
                yield from walk(child)

        with patch("automations.bordero.panel.SqlConferenceProvider") as source:
            source.return_value.fetch.return_value = (
                map_document(document()),
                map_document(document(iddocumentocapcarbase=124, formadepagamento="Boleto")),
            )
            web_download = MagicMock(return_value="/downloads/test-token")
            main(page, web_download=web_download)
            button = next(item for root in controls for item in walk(root) if isinstance(item, ft.ElevatedButton) and item.text.startswith("Consultar"))
            asyncio.run(button.on_click(None))
            source.return_value.fetch.assert_called_once()
            payment_filter = next(item for root in controls for item in walk(root) if isinstance(item, ft.Dropdown) and item.label == "Forma de pagamento")
            self.assertEqual(len(payment_filter.options), 3)
            payment_filter.value = "method:Transferência"
            payment_filter.on_change(None)
            tables = [item for root in controls for item in walk(root) if isinstance(item, ft.DataTable)]
            self.assertEqual(len(tables[0].rows), 1)
            headers = [column.label.value for column in tables[0].columns]
            self.assertIn("CPF/CNPJ", headers)
            self.assertEqual(headers[-1], "Histórico")
            review_button = tables[0].rows[0].cells[headers.index("Conferir")].content
            review_button.on_click(None)
            dialog = page.open.call_args.args[0]
            self.assertIsInstance(dialog, ft.AlertDialog)
            self.assertEqual(dialog.title.value, "Conferir pagamento")
            decision = next(item for item in walk(dialog.content) if isinstance(item, ft.Dropdown) and item.label == "Classificação")
            decision.value = "Conferido"
            dialog.actions[-1].on_click(None)
            page.close.assert_called_once_with(dialog)
            tables = [item for root in controls for item in walk(root) if isinstance(item, ft.DataTable)]
            self.assertEqual(tables[0].rows[0].cells[headers.index("Conferência")].content.value, "Conferido")
            download_button = next(item for root in controls for item in walk(root) if isinstance(item, ft.OutlinedButton) and item.text == "Baixar Excel")
            download_button.on_click(None)
            web_download.assert_called_once()
            exported_rows = web_download.call_args.args[0]
            self.assertEqual(len(exported_rows), 1)
            self.assertEqual(exported_rows[0].situation, "Conferido")
            page.launch_url.assert_called_once_with("/downloads/test-token", web_window_name="_self")
            asyncio.run(button.on_click(None))
            tables = [item for root in controls for item in walk(root) if isinstance(item, ft.DataTable)]
            self.assertEqual(tables[0].rows[0].cells[headers.index("Conferência")].content.value, "A conferir")
            logo = next(item for root in controls for item in walk(root) if isinstance(item, ft.Image) and item.src == "/carmel-logo.png")
            self.assertEqual(logo.semantics_label, "Carmel Hotéis")
            theme_button = next(item for root in controls for item in walk(root) if isinstance(item, ft.IconButton) and item.tooltip == "Usar modo claro")
            theme_button.on_click(None)
            self.assertEqual(page.bgcolor, "#D8D2C8")
            self.assertEqual(payment_filter.bgcolor, "#EFEBE4")
            self.assertEqual(payment_filter.color, "#151311")
            self.assertEqual(payment_filter.text_style.weight, ft.FontWeight.W_500)
            light_table = next(item for root in controls for item in walk(root) if isinstance(item, ft.DataTable))
            self.assertEqual(light_table.rows[0].cells[0].content.weight, ft.FontWeight.W_500)
            self.assertEqual(light_table.rows[0].cells[0].content.color, "#151311")
            self.assertEqual(payment_filter.hover_color, "#DBD4CA")
            self.assertEqual(page.theme.color_scheme.surface_container_high, "#EFEBE4")
            status_icon = next(item for root in controls for item in walk(root) if isinstance(item, ft.Icon) and item.name == ft.Icons.CHECK_CIRCLE_OUTLINE)
            self.assertEqual(status_icon.color, "#205E3F")
            self.assertEqual(download_button.icon_color, "#205E3F")
            theme_button.on_click(None)
            self.assertEqual(page.bgcolor, "#1D1D1B")
            self.assertEqual(payment_filter.bgcolor, "#292623")
            self.assertEqual(payment_filter.color, "#E8E7E1")
            self.assertEqual(status_icon.color, "#83C79B")
            self.assertEqual(download_button.icon_color, "#83C79B")

    def test_maps_missing_fields_without_approving(self):
        row = map_document(document())
        self.assertEqual(row.situation, "A conferir")
        self.assertEqual(row.scheduled_date, None)
        self.assertEqual(row.balance, Decimal("100.50"))
        self.assertEqual(row.supplier, "Fornecedor exemplo")
        self.assertNotIn("iddocumentocapcarbase", row.details)
        self.assertEqual(row.responsible, "Ana")
        self.assertEqual(row.source_situation, "Aprovado")

    def test_attachment_count_preserves_unknown_and_presence(self):
        self.assertEqual(map_document(document(numerodeanexos=None)).attachment_summary, "Não informado")
        self.assertEqual(map_document(document(numerodeanexos=0)).attachment_summary, "Sem anexo")
        self.assertEqual(map_document(document(numerodeanexos=2)).attachment_summary, "Sim (2)")

    def test_payment_and_review_filters_are_combined(self):
        transfer = map_document(document())
        boleto = replace(map_document(document(formadepagamento="Boleto")), situation="Pendente")
        missing = map_document(document(formadepagamento=None))
        rows = (transfer, boleto, missing)
        self.assertEqual(filter_conferences(rows, payment_method="Boleto"), (boleto,))
        self.assertEqual(filter_conferences(rows, "Conferido", "Boleto"), ())
        self.assertEqual(filter_conferences(rows, "Pendente", "Boleto"), (boleto,))
        self.assertEqual(filter_conferences(rows, payment_method=""), (missing,))
        self.assertEqual(filter_conferences(rows), rows)

    def test_filter_is_parameterized_and_last_day_inclusive(self):
        connection = MagicMock()
        cursor = connection.__enter__.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchall.return_value = [document()]
        start, end = date(2026, 10, 1), date(2026, 10, 5)
        with patch("automations.bordero.database.connection_settings", return_value={}), patch("automations.bordero.database.psycopg.connect", return_value=connection):
            rows = SqlConferenceProvider().fetch(start, end, "MAGNA' OR 1=1 --")
        query, params = cursor.execute.call_args.args
        self.assertIn("cpfcnpj", query)
        self.assertNotIn("OR 1=1", query)
        self.assertEqual(params, [start, end + timedelta(days=1), "MAGNA' OR 1=1 --"])
        self.assertTrue(connection.__enter__.return_value.read_only)
        self.assertEqual(len(rows), 1)

    def test_classification_is_in_memory_and_requires_reason(self):
        row = map_document(document())
        with self.assertRaisesRegex(ValueError, "motivo"):
            classify_conference(row, "Pendente")
        pending = classify_conference(row, "Pendente", "Falta comprovante")
        self.assertEqual(pending.situation, "Pendente")
        self.assertEqual(pending.pending_reason, "Falta comprovante")
        self.assertEqual(row.situation, "A conferir")
        self.assertEqual(classify_conference(pending, "Conferido").pending_reason, "")
