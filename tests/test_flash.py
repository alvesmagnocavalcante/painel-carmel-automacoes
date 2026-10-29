"""Tests for Flash corporate card reconciliation."""

import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import MagicMock

import flet as ft
from openpyxl import Workbook, load_workbook

from automations.cartao_flash.data import (
    STATUS_PENDING_LAUNCH,
    STATUS_PENDING_SETTLEMENT,
    STATUS_SETTLED,
    CmflexFlashDocument,
    FlashReconciliationItem,
    FlashStatement,
    FlashTransaction,
    merge_flash_statements,
    normalize_text,
    parse_flash_amount,
    parse_flash_csv,
    parse_flash_date,
    parse_flash_statement,
    parse_flash_xlsx,
)
from automations.cartao_flash.export import export_flash_reconciliation, generate_email_text
from automations.cartao_flash.panel import main as build_flash_panel
from automations.cartao_flash.reconciliation import reconcile_flash_statement
from automations.cartao_flash.store import FlashStore


SAMPLE_CSV = """Data,Movimentação,Valor,Pessoa,Pagamento,Prestação de contas
21/10/2026 09:22,MP*ROBINSONGASES Osasco BRA Conveniência,"- R$ 200,00",DAVID DE OLIVEIRA BARROS,Cartão físico - Conta final 402,Pendente
20/10/2026 17:35,PAYGO*NORDESTE DIST Eus bio BRA Conveniência,"- R$ 157,50",DAVID DE OLIVEIRA BARROS,Cartão físico - Conta final 402,Pendente
14/10/2026 14:34,DLKNET *AC EUSEBIO EUSEBIO BRA Entrega e frete,"- R$ 48,40",Andréa De Boni Nottingham,Cartão físico - Conta final 986,Pendente
09/10/2026 09:59,Depósito único,"+ R$ 90,90",Andréa De Boni Nottingham,Carteira corporativa - Agendado por LAIS EDUARDA DO NASCIMENTO MENDONÇA,-
31/08/2026 16:43,MAXLIMP FORTALEZA BRA Conveniência,"- R$ 100,00",ALYNE SOUZA ROCHA,Cartão físico - Conta final 742,Pendente
21/08/2026 15:21,JAGUAR TINTAS FORTALEZA BRA Conveniência,"- R$ 85,40",ALYNE SOUZA ROCHA,Cartão físico - Conta final 742,Pendente
"""


class FlashDataTests(TestCase):
    def test_amount_and_date_parsing(self):
        amt, is_credit = parse_flash_amount("- R$ 200,00")
        self.assertEqual(amt, Decimal("200.00"))
        self.assertFalse(is_credit)

        amt_dep, is_credit_dep = parse_flash_amount("+ R$ 90,90")
        self.assertEqual(amt_dep, Decimal("90.90"))
        self.assertTrue(is_credit_dep)

        dt = parse_flash_date("21/10/2026 09:22")
        self.assertEqual(dt, date(2026, 10, 21))

    def test_display_batch_property(self):
        doc_with_batch = CmflexFlashDocument(
            document_id=1,
            company_id=1,
            company="CARMEL",
            launch_number="100",
            supplier="POSTO",
            tax_id="",
            launch_date=date(2026, 10, 20),
            due_date=None,
            scheduled_date=None,
            amount=Decimal("50.00"),
            balance=Decimal("0.00"),
            document_status=2,
            payment_method="",
            holder="",
            batch_id="45210",
        )
        self.assertEqual(doc_with_batch.display_batch, "Lote #45210")

        doc_without_batch = CmflexFlashDocument(
            document_id=2,
            company_id=1,
            company="CARMEL",
            launch_number="101",
            supplier="POSTO",
            tax_id="",
            launch_date=date(2026, 10, 20),
            due_date=None,
            scheduled_date=None,
            amount=Decimal("50.00"),
            balance=Decimal("0.00"),
            document_status=2,
            payment_method="",
            holder="",
            batch_id="",
        )
        self.assertEqual(doc_without_batch.display_batch, "Sem Lote")

    def test_parse_csv_statement(self):
        statement = parse_flash_csv(SAMPLE_CSV)
        self.assertEqual(len(statement.transactions), 6)
        expenses = statement.expense_transactions
        self.assertEqual(len(expenses), 5)
        self.assertEqual(statement.total_deposits, Decimal("90.90"))
        self.assertEqual(statement.total_expenses, Decimal("591.30"))

        first = expenses[0]
        self.assertEqual(first.person, "DAVID DE OLIVEIRA BARROS")
        self.assertEqual(first.amount, Decimal("200.00"))
        self.assertEqual(first.month_name, "OUTUBRO")

    def test_parse_xlsx_statement(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Data", "Movimentação", "Valor", "Pessoa", "Pagamento", "Prestação de contas"])
        ws.append(["21/08/2026 15:21", "JAGUAR TINTAS", "-85.40", "ALYNE SOUZA", "Cartão físico", "Pendente"])
        buf = BytesIO()
        wb.save(buf)

        statement = parse_flash_xlsx(buf.getvalue())
        self.assertEqual(len(statement.transactions), 1)
        tx = statement.transactions[0]
        self.assertEqual(tx.amount, Decimal("85.40"))
        self.assertEqual(tx.person, "ALYNE SOUZA")

    def test_merge_flash_statements_combines_and_deduplicates(self):
        tx1 = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 10, 20),
            raw_date="20/10/2026 10:00",
            description="HOTEL POSTO",
            amount=Decimal("150.00"),
            person="MARIA SILVA",
            payment_method="Físico",
        )
        tx2 = FlashTransaction(
            row_index=2,
            purchase_date=date(2026, 10, 21),
            raw_date="21/10/2026 12:00",
            description="RESTAURANTE",
            amount=Decimal("80.00"),
            person="MARIA SILVA",
            payment_method="Físico",
        )
        # Duplicate of tx1 from another file
        tx1_dup = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 10, 20),
            raw_date="20/10/2026 10:00",
            description="HOTEL POSTO",
            amount=Decimal("150.00"),
            person="MARIA SILVA",
            payment_method="Físico",
        )
        tx3 = FlashTransaction(
            row_index=2,
            purchase_date=date(2026, 10, 22),
            raw_date="22/10/2026 15:00",
            description="LIVRARIA",
            amount=Decimal("50.00"),
            person="JOAO SOUZA",
            payment_method="Virtual",
        )

        st1 = FlashStatement(transactions=(tx1, tx2), file_name="outubro_1.xlsx")
        st2 = FlashStatement(transactions=(tx1_dup, tx3), file_name="outubro_2.xlsx")

        merged = merge_flash_statements([st1, st2])
        self.assertEqual(len(merged.transactions), 3)
        self.assertEqual(merged.total_expenses, Decimal("280.00"))
        # Should be ordered descending by purchase date
        self.assertEqual(merged.transactions[0].description, "LIVRARIA")
        self.assertEqual(merged.transactions[1].description, "RESTAURANTE")
        self.assertEqual(merged.transactions[2].description, "HOTEL POSTO")
        # Row index correctly reindexed 1, 2, 3
        self.assertEqual([t.row_index for t in merged.transactions], [1, 2, 3])


class FlashReconciliationTests(TestCase):
    def test_reconciliation_identifies_all_three_statuses(self):
        tx_settled = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 8, 21),
            raw_date="21/08/2026",
            description="JAGUAR TINTAS FORTALEZA BRA",
            amount=Decimal("85.40"),
            person="ALYNE SOUZA ROCHA",
            hotel="MAGNA PRAIA",
        )
        tx_pending_prestacao = FlashTransaction(
            row_index=2,
            purchase_date=date(2026, 10, 27),
            raw_date="27/10/2026",
            description="T R F DE OLIVEIRA",
            amount=Decimal("67.00"),
            person="DAVID DE OLIVEIRA",
            hotel="MAGNA PRAIA",
        )
        tx_pending_lancamento = FlashTransaction(
            row_index=3,
            purchase_date=date(2026, 8, 31),
            raw_date="31/08/2026",
            description="MAXLIMP FORTALEZA BRA",
            amount=Decimal("100.00"),
            person="ALYNE SOUZA ROCHA",
            hotel="MAGNA PRAIA",
        )

        doc_settled = CmflexFlashDocument(
            document_id=1,
            company_id=4,
            company="MAGNA PRAIA",
            launch_number="333534",
            supplier="JAGUAR COMERCIO DE TINTAS LTDA",
            tax_id="07221472000166",
            launch_date=date(2026, 8, 28),
            due_date=date(2026, 9, 9),
            scheduled_date=date(2026, 9, 9),
            amount=Decimal("85.40"),
            balance=Decimal("0.00"),
            document_status=2,
            payment_method="Cartão Fundo Fixo",
            holder="FUNDO FIXO MAGNA",
        )
        doc_pending = CmflexFlashDocument(
            document_id=2,
            company_id=4,
            company="MAGNA PRAIA",
            launch_number="722",
            supplier="T R F DE OLIVEIRA",
            tax_id="04000842000102",
            launch_date=date(2026, 10, 27),
            due_date=date(2026, 11, 6),
            scheduled_date=date(2026, 11, 6),
            amount=Decimal("67.00"),
            balance=Decimal("67.00"),
            document_status=0,
            payment_method="Cartão Fundo Fixo",
            holder="FUNDO FIXO MAGNA",
        )

        items = reconcile_flash_statement(
            (tx_settled, tx_pending_prestacao, tx_pending_lancamento),
            (doc_settled, doc_pending),
        )

        self.assertEqual(len(items), 3)
        self.assertEqual(items[0].status, STATUS_SETTLED)
        self.assertEqual(items[0].matched_document, doc_settled)

        self.assertEqual(items[1].status, STATUS_PENDING_SETTLEMENT)
        self.assertEqual(items[1].matched_document, doc_pending)

        self.assertEqual(items[2].status, STATUS_PENDING_LAUNCH)
        self.assertIsNone(items[2].matched_document)

    def test_same_amount_and_date_without_supplier_evidence_is_not_a_match(self):
        transaction = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 10, 14),
            raw_date="14/10/2026",
            description="DLKNET AC EUSEBIO",
            amount=Decimal("48.40"),
            person="ANDREA NOTTINGHAM",
        )
        unrelated = CmflexFlashDocument(
            document_id=99,
            company_id=1,
            company="CHARME HOSPEDAGEM",
            launch_number="921967615",
            supplier="EMPRESA BRASILEIRA DE CORREIOS E TELEGRAFOS",
            tax_id="",
            launch_date=date(2026, 10, 14),
            due_date=None,
            scheduled_date=None,
            amount=Decimal("48.40"),
            balance=Decimal("0.00"),
            document_status=2,
            payment_method="Cartão Fundo Fixo",
            holder="",
        )

        item = reconcile_flash_statement((transaction,), (unrelated,))[0]

        self.assertEqual(item.status, STATUS_PENDING_LAUNCH)
        self.assertIsNone(item.matched_document)

    def test_contained_merchant_name_can_match_supplier(self):
        transaction = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 10, 21),
            raw_date="21/10/2026",
            description="MP ROBINSONGASES OSASCO",
            amount=Decimal("200.00"),
            person="DAVID BARROS",
        )
        document = CmflexFlashDocument(
            document_id=100,
            company_id=1,
            company="CHARME HOSPEDAGEM",
            launch_number="3739",
            supplier="PAULO ROBINSON COSTA SILVA ME",
            tax_id="",
            launch_date=date(2026, 10, 21),
            due_date=None,
            scheduled_date=None,
            amount=Decimal("200.00"),
            balance=Decimal("200.00"),
            document_status=0,
            payment_method="Cartão Fundo Fixo",
            holder="",
        )

        item = reconcile_flash_statement((transaction,), (document,))[0]

        self.assertEqual(item.matched_document, document)
        self.assertEqual(item.status, STATUS_PENDING_SETTLEMENT)


class FlashStoreTests(TestCase):
    def test_sqlite_store_persists_review_notes(self):
        with TemporaryDirectory() as tmp_dir:
            db_path = Path(tmp_dir) / "flash_test.sqlite3"
            store = FlashStore(db_path)

            tx = FlashTransaction(
                row_index=1,
                purchase_date=date(2026, 10, 21),
                raw_date="21/10/2026",
                description="ROBINSONGASES",
                amount=Decimal("200.00"),
                person="DAVID DE OLIVEIRA",
            )
            store.save_review(tx, status=STATUS_PENDING_LAUNCH, notes="Aguardando comprovante com a gerência")

            # Re-read
            store2 = FlashStore(db_path)
            reviews = store2.load_reviews()
            self.assertEqual(len(reviews), 1)
            key = next(iter(reviews))
            self.assertEqual(reviews[key], (STATUS_PENDING_LAUNCH, "Aguardando comprovante com a gerência"))


class FlashExportTests(TestCase):
    def test_excel_export_creates_valid_workbook(self):
        tx = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 8, 31),
            raw_date="31/08/2026",
            description="MAXLIMP FORTALEZA",
            amount=Decimal("100.00"),
            person="ALYNE SOUZA ROCHA",
            hotel="MAGNA PRAIA",
        )
        item = FlashReconciliationItem(transaction=tx, status=STATUS_PENDING_LAUNCH, hotel="MAGNA PRAIA")
        doc = CmflexFlashDocument(
            document_id=1,
            company_id=1,
            company="CARMEL",
            launch_number="100",
            supplier="MAXLIMP",
            tax_id="",
            launch_date=date(2026, 8, 31),
            due_date=None,
            scheduled_date=None,
            amount=Decimal("100.00"),
            balance=Decimal("0.00"),
            document_status=2,
            payment_method="",
            holder="",
            batch_id="9988",
        )
        item_matched = FlashReconciliationItem(
            transaction=tx,
            status=STATUS_SETTLED,
            hotel="MAGNA PRAIA",
            matched_document=doc,
        )
        data = export_flash_reconciliation((item, item_matched))
        self.assertIsInstance(data, bytes)
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=9).value, "LOTE CMFLEX")
        self.assertEqual(ws.cell(row=2, column=1).value, "AGOSTO")
        self.assertEqual(ws.cell(row=2, column=3).value, "MAXLIMP FORTALEZA")
        self.assertEqual(ws.cell(row=2, column=6).value, STATUS_PENDING_LAUNCH)
        self.assertFalse(ws.cell(row=2, column=9).value)
        self.assertEqual(ws.cell(row=3, column=9).value, "Lote #9988")

    def test_email_text_generation(self):
        tx = FlashTransaction(
            row_index=1,
            purchase_date=date(2026, 8, 31),
            raw_date="31/08/2026",
            description="MAXLIMP",
            amount=Decimal("100.00"),
            person="ALYNE SOUZA ROCHA",
            hotel="MAGNA PRAIA",
        )
        item = FlashReconciliationItem(transaction=tx, status=STATUS_PENDING_LAUNCH, hotel="MAGNA PRAIA")
        text = generate_email_text((item,))
        self.assertIn("MAGNA PRAIA", text)
        self.assertIn("ALYNE SOUZA ROCHA", text)
        self.assertIn("MAXLIMP", text)
        self.assertIn("PENDENTE DE LANÇAMENTO", text)


class FlashPanelUiTests(TestCase):
    @staticmethod
    def _walk(control):
        yield control
        content = getattr(control, "content", None)
        if isinstance(content, ft.Control):
            yield from FlashPanelUiTests._walk(content)
        for child in getattr(control, "controls", ()) or ():
            yield from FlashPanelUiTests._walk(child)

    def test_panel_builds_and_loads_sample(self):
        page = MagicMock()
        page.web = True
        page.overlay = []

        with TemporaryDirectory() as tmp_dir:
            store = FlashStore(Path(tmp_dir) / "flash.sqlite3")
            mock_provider = MagicMock()
            mock_provider.fetch_documents.return_value = ()

            panel_view = build_flash_panel(
                page,
                download_file=MagicMock(),
                provider=mock_provider,
                upload_dir=tmp_dir,
                store=store,
            )
            self.assertIsInstance(panel_view.control, ft.Column)
            self.assertEqual(panel_view.control.scroll, ft.ScrollMode.AUTO)
            texts = [
                item.value
                for item in self._walk(panel_view.control)
                if isinstance(item, ft.Text)
            ]
            self.assertIn("Nenhum extrato do Flash carregado", texts)
            empty_state = next(
                item
                for item in self._walk(panel_view.control)
                if isinstance(item, ft.Container)
                and item.height == 320
                and any(
                    isinstance(child, ft.Text)
                    and child.value == "Nenhum extrato do Flash carregado"
                    for child in self._walk(item)
                )
            )
            self.assertEqual(empty_state.height, 320)
            panel_view.apply_theme(False)
            panel_view.apply_theme(True)

    def test_upload_queries_only_reconciliation_period_without_blocking_callback(self):
        page = MagicMock(web=True)
        page.overlay = []
        page.get_upload_url.return_value = "/upload?signature=test"
        page.run_task.side_effect = lambda handler, *args: asyncio.run(handler(*args))

        with TemporaryDirectory() as tmp_dir:
            store = MagicMock()
            store.load_reviews.return_value = {}
            provider = MagicMock()
            provider.fetch_documents.return_value = ()
            panel_view = build_flash_panel(
                page,
                download_file=MagicMock(),
                provider=provider,
                upload_dir=tmp_dir,
                store=store,
                available_width=lambda: 1900,
            )
            picker = page.overlay[0]
            picker.upload = MagicMock()
            picker.on_result(MagicMock(files=[SimpleNamespace(name="flash.csv", size=len(SAMPLE_CSV))]))
            token = page.get_upload_url.call_args.args[0]
            uploaded = Path(tmp_dir) / token
            uploaded.write_text(SAMPLE_CSV, encoding="utf-8")

            picker.on_upload(MagicMock(error=None, progress=1.0))

            provider.fetch_documents.assert_called_once_with(
                start_date=date(2026, 7, 22),
                end_date=date(2026, 11, 21),
            )
            self.assertFalse(uploaded.exists())
            self.assertTrue(panel_view.control._build_add_commands())
            self.assertFalse(
                any(
                    isinstance(item, ft.Row)
                    and item.vertical_alignment == ft.CrossAxisAlignment.BASELINE
                    for item in self._walk(panel_view.control)
                )
            )
            table = next(
                item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable)
            )
            self.assertEqual(table.width, 1900)
            self.assertGreater(table.rows[0].cells[1].content.width, 230)
            self.assertTrue(
                all(
                    isinstance(cell.content, (ft.Text, ft.IconButton, ft.Container))
                    for row in table.rows
                    for cell in row.cells
                )
            )
            wrapped_cells = [
                cell.content.content
                for row in table.rows
                for cell in row.cells
                if isinstance(cell.content, ft.Container)
                and isinstance(cell.content.content, ft.Text)
            ]
            self.assertTrue(any(text.max_lines and text.max_lines > 1 for text in wrapped_cells))
            viewport = next(
                item
                for item in self._walk(panel_view.control)
                if isinstance(item, ft.Row)
                and item.scroll == ft.ScrollMode.ALWAYS
                and table in tuple(self._walk(item))
            )
            self.assertIsInstance(viewport.height, (int, float))

    def test_invalid_upload_is_removed_and_shows_feedback(self):
        page = MagicMock(web=True)
        page.overlay = []
        page.get_upload_url.return_value = "/upload?signature=test"

        with TemporaryDirectory() as tmp_dir:
            panel_view = build_flash_panel(
                page,
                provider=MagicMock(),
                upload_dir=tmp_dir,
                store=MagicMock(),
            )
            picker = page.overlay[0]
            picker.upload = MagicMock()
            picker.on_result(MagicMock(files=[SimpleNamespace(name="flash.csv", size=7)]))
            uploaded = Path(tmp_dir) / page.get_upload_url.call_args.args[0]
            uploaded.write_bytes(b"invalid")

            picker.on_upload(MagicMock(error=None, progress=1.0))

            self.assertFalse(uploaded.exists())
            self.assertTrue(page.snack_bar.open)
            self.assertTrue(panel_view.control._build_add_commands())

    def test_multi_file_upload_merges_all_transactions(self):
        page = MagicMock(web=True)
        page.overlay = []
        page.get_upload_url.side_effect = lambda token, *args: f"/upload?token={token}"
        page.run_task.side_effect = lambda handler, *args: asyncio.run(handler(*args))

        csv_file1 = "Data,Movimentação,Valor,Pessoa,Pagamento,Prestação\n21/10/2026 10:00,LOJA A,-100.00,MARIA,Físico,Pendente\n"
        csv_file2 = "Data,Movimentação,Valor,Pessoa,Pagamento,Prestação\n22/10/2026 11:00,LOJA B,-200.00,JOSE,Virtual,Pendente\n"

        with TemporaryDirectory() as tmp_dir:
            store = MagicMock()
            store.load_reviews.return_value = {}
            provider = MagicMock()
            provider.fetch_documents.return_value = ()

            panel_view = build_flash_panel(
                page,
                download_file=MagicMock(),
                provider=provider,
                upload_dir=tmp_dir,
                store=store,
            )
            picker = page.overlay[0]
            picker.upload = MagicMock()

            # Select 2 files at once
            picker.on_result(
                MagicMock(
                    files=[
                        SimpleNamespace(name="fatura_a.csv", size=len(csv_file1)),
                        SimpleNamespace(name="fatura_b.csv", size=len(csv_file2)),
                    ]
                )
            )

            # Two upload calls were made
            call_tokens = [call.args[0] for call in page.get_upload_url.call_args_list]
            self.assertEqual(len(call_tokens), 2)

            (Path(tmp_dir) / call_tokens[0]).write_text(csv_file1, encoding="utf-8")
            (Path(tmp_dir) / call_tokens[1]).write_text(csv_file2, encoding="utf-8")

            # Progress events for each file
            picker.on_upload(MagicMock(file_name="fatura_a.csv", progress=1.0, error=None))
            # First file done, second not yet - CMFlex not fetched yet
            self.assertEqual(provider.fetch_documents.call_count, 0)

            # Second file completes
            picker.on_upload(MagicMock(file_name="fatura_b.csv", progress=1.0, error=None))

            # Both files parsed, merged, and CMFlex queried
            provider.fetch_documents.assert_called_once()
            table = next(
                item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable)
            )
            self.assertEqual(len(table.rows), 2)
            self.assertTrue(panel_view.control._build_add_commands())
