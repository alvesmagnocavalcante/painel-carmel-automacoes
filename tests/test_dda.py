"""Unit tests for DDA reconciliation module."""

import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import MagicMock

import flet as ft
from openpyxl import Workbook, load_workbook

from automations.dda.data import (
    OBS_NAO_LOCALIZADO,
    OBS_SEM_ALTERACAO,
    OBS_SEM_ANEXO_E_FORMA,
    OBS_SEM_BOLETO,
    OBS_SEM_BOLETO_E_DADOS_BAN,
    OBS_VALOR_DIVERGENTE,
    STATUS_LANC,
    STATUS_LANC_OC,
    STATUS_NAO_LANC,
    CmflexOpenDocument,
    DdaReconciliationItem,
    DdaTransaction,
    clean_tax_id,
    format_tax_id,
    normalize_text,
    parse_amount,
    parse_cmflex_export_workbook,
    parse_date,
    parse_dda_workbook,
)
from automations.dda.export import export_dda_reconciliation, generate_collection_email
from automations.dda.panel import main as build_dda_panel
from automations.dda.reconciliation import reconcile_dda
from automations.dda.store import DdaStore, make_transaction_key


class DdaDataTests(TestCase):
    def test_amount_parsing(self):
        self.assertEqual(parse_amount("36.869,20"), Decimal("36869.20"))
        self.assertEqual(parse_amount("R$ 21.840,37"), Decimal("21840.37"))
        self.assertEqual(parse_amount("850"), Decimal("850.00"))
        self.assertEqual(parse_amount(1234.56), Decimal("1234.56"))
        self.assertEqual(parse_amount(None), Decimal("0.00"))
        self.assertEqual(parse_amount("invalid"), Decimal("0.00"))

    def test_date_parsing(self):
        self.assertEqual(parse_date("25/07/2026"), date(2026, 7, 25))
        self.assertEqual(parse_date("2026-07-25"), date(2026, 7, 25))
        self.assertEqual(parse_date(date(2026, 7, 25)), date(2026, 7, 25))
        self.assertIsNone(parse_date(None))
        self.assertIsNone(parse_date(""))

    def test_tax_id_helpers(self):
        self.assertEqual(clean_tax_id("07.647.045/0001-50"), "07647045000150")
        self.assertEqual(format_tax_id("07647045000150"), "07.647.045/0001-50")
        self.assertEqual(format_tax_id("12345678901"), "123.456.789-01")

    def test_parse_dda_workbook_multi_sheet(self):
        wb = Workbook()
        # Sheet 1: TAIBA
        ws1 = wb.active
        ws1.title = "TAIBA"
        ws1.append([
            "EMPRESA", "BANCO", "CNPJ", "FORNECEDOR", "DOCUMENTO", "TIPO", "SITUAÇÃO", "DATA EMISSÃO", "DATA VENCIMENTO", "VALOR"
        ])
        ws1.append([
            "CARMEL TAÍBA", "341", "07.647.045/0001-50", "FORNECEDOR ALPHA", "1001", "Cobrança", "A pagar", "20/07/2026", "25/07/2026", "1.500,00"
        ])

        # Sheet 2: WIND
        ws2 = wb.create_sheet(title="WIND")
        ws2.append([
            "EMPRESA", "BANCO", "CNPJ", "FORNECEDOR", "DOCUMENTO", "TIPO", "SITUAÇÃO", "DATA EMISSÃO", "DATA VENCIMENTO", "VALOR"
        ])
        ws2.append([
            "CARMEL WIND", "237", "09.370.323/0001-41", "FORNECEDOR BETA", "2002", "Cobrança", "A pagar", "22/07/2026", "28/07/2026", "2.800,50"
        ])

        buf = BytesIO()
        wb.save(buf)

        txs = parse_dda_workbook(buf.getvalue())
        self.assertEqual(len(txs), 2)
        self.assertEqual(txs[0].sheet_name, "TAIBA")
        self.assertEqual(txs[0].amount, Decimal("1500.00"))
        self.assertEqual(txs[0].supplier, "FORNECEDOR ALPHA")
        self.assertEqual(txs[0].clean_cnpj, "07647045000150")

        self.assertEqual(txs[1].sheet_name, "WIND")
        self.assertEqual(txs[1].amount, Decimal("2800.50"))
        self.assertEqual(txs[1].supplier, "FORNECEDOR BETA")

    def test_parse_dda_csv_itau_format(self):
        csv_content = """
,,,,,,,,,,
Dados da conta,,,,,,,,,,
Nome da empresa: MAGNA PRAIA HOTEL LTDA,,,,,,,,,,
Agência/Conta: 1646 / 0069024-6,,,,,,,,,,
CPF/CNPJ:  02.333.096/0001-35,,,,,,,,,,
Boletos,,,,,,,,,,
Data/Hora: 10/09/2026 às 16:45:24,,,,,,,,,,
Pagador/Agregado,Beneficiário,CPF/CNPJ,Venc.,Nº Doc.,Até Venc.,A Pagar,Tipo de Boleto,Banco,Observações,Código de barras
CM CENTRAL DE SERVICOS ADMINISTRATIVOS L          ,APICE SERVICE SOLUCOES LTDA                       ,30.013.192/0001-31,16/09/2026,901008589,"R$ 794,80","R$ 794,80",DM Duplicata Mercantil,077,Este boleto sofreu alteração.,07799157100000794800001112071892090841278321
MAGNA PRAIA HOTEL LTDA                            ,TRAMONTINA NORDESTE S/A                           ,15.205.628/0001-09,16/09/2026,0938028-00,"R$ 801,39","R$ 801,39",DM Duplicata Mercantil,001,,00191157100000801390628462414134290070993111
CARMEL WIND RESORT LTDA                           ,BANCO ITAU S/A                                    ,60.701.190/0001-04,16/09/2026,00500474 1,"R$ 8.969,33","R$ 8.969,33",DM Duplicata Mercantil,341,Instrução,34196157100008969331090024649168144277368000
CARMEL TAIBA EXCLUSIVE RESORT                     ,ECOLAV SERVICOS L LTDA                            ,24.084.433/0001-30,16/09/2026,3616,"R$ 15.211,28","R$ 15.211,28",DS Duplicata de Serviço,341,,34197157100015211281570004131138789409466000
"""
        txs = parse_dda_workbook = None  # test CSV parser directly
        from automations.dda.data import parse_dda_csv, parse_dda_file
        txs = parse_dda_csv(csv_content, file_name="boletos_itau.csv")
        self.assertEqual(len(txs), 4)

        # CM Unit
        self.assertEqual(txs[0].sheet_name, "CM")
        self.assertEqual(txs[0].supplier, "APICE SERVICE SOLUCOES LTDA")
        self.assertEqual(txs[0].clean_cnpj, "30013192000131")
        self.assertEqual(txs[0].amount, Decimal("794.80"))
        self.assertEqual(txs[0].bank_code, "077")
        self.assertEqual(txs[0].due_date, date(2026, 9, 16))

        # Magna Unit
        self.assertEqual(txs[1].sheet_name, "MAGNA")
        self.assertEqual(txs[1].amount, Decimal("801.39"))

        # Wind Unit
        self.assertEqual(txs[2].sheet_name, "WIND")
        self.assertEqual(txs[2].amount, Decimal("8969.33"))

        # Taiba Unit
        self.assertEqual(txs[3].sheet_name, "TAIBA")
        self.assertEqual(txs[3].amount, Decimal("15211.28"))

        # Auto detect with parse_dda_file
        txs_auto = parse_dda_file(csv_content.encode("utf-8"), file_name="dda.csv")
        self.assertEqual(len(txs_auto), 4)

    def test_parse_cmflex_export_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "PagamentosEmAbertoDTO"
        ws.append([
            "Tabela", "Pai", "IdPai", "FornecedorNomeFantasia", "FornecedorRazaoSocial", "FornecedorDocumento",
            "StatusBPM", "IdInstanciaBPM", "FormaDef", "Portador/Forma", "Numero", "Complemento", "Empresa", "Saldo",
            "SaldoOutraMoeda", "DataLanc", "DataVenc", "DataProg", "Historico", "Responsavel"
        ])
        ws.append([
            "Pagamento", "NULL", "NULL", "FORNECEDOR ALPHA", "FORNECEDOR ALPHA LTDA", "07647045000150",
            "NULL", "0", "NULL", "Boletos de BANCO ITAU", "1001", "1", "CARMEL TAÍBA", "1500,00",
            "0", "20/07/2026", "25/07/2026", "25/07/2026", "Pagamento de NF 1001", "FRANCISCO LUAN"
        ])
        buf = BytesIO()
        wb.save(buf)

        docs = parse_cmflex_export_workbook(buf.getvalue())
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].supplier, "FORNECEDOR ALPHA")
        self.assertEqual(docs[0].amount, Decimal("1500.00"))
        self.assertEqual(docs[0].responsible, "FRANCISCO LUAN")


class DdaReconciliationTests(TestCase):
    def setUp(self):
        self.tx1 = DdaTransaction(
            sheet_name="TAIBA",
            row_index=2,
            company="CARMEL TAÍBA",
            bank_code="341",
            tax_id="07.647.045/0001-50",
            supplier="FORNECEDOR ALPHA COMERCIO",
            document_number="1001",
            charge_type="Cobrança",
            situation="A pagar",
            issue_date=date(2026, 7, 20),
            due_date=date(2026, 7, 25),
            amount=Decimal("1500.00"),
        )
        self.tx2 = DdaTransaction(
            sheet_name="TAIBA",
            row_index=3,
            company="CARMEL TAÍBA",
            bank_code="341",
            tax_id="09.370.323/0001-41",
            supplier="MANDACARU SERVICOS",
            document_number="2002",
            charge_type="Cobrança",
            situation="A pagar",
            issue_date=date(2026, 7, 20),
            due_date=date(2026, 7, 26),
            amount=Decimal("2000.00"),
        )
        self.tx_unmatched = DdaTransaction(
            sheet_name="WIND",
            row_index=4,
            company="CARMEL WIND",
            bank_code="237",
            tax_id="99.999.999/0001-99",
            supplier="FORNECEDOR DESCONHECIDO",
            document_number="9999",
            charge_type="Cobrança",
            situation="A pagar",
            issue_date=date(2026, 7, 20),
            due_date=date(2026, 7, 25),
            amount=Decimal("750.00"),
        )

    def test_reconciliation_matches_and_detects_boleto(self):
        doc1 = CmflexOpenDocument(
            document_id=1,
            company_id=1,
            company="CARMEL TAÍBA",
            launch_number="1001",
            installment="1",
            supplier="FORNECEDOR ALPHA COMERCIO LTDA",
            tax_id="07647045000150",
            reference="NF 1001",
            amount=Decimal("1500.00"),
            balance=Decimal("1500.00"),
            launch_date=date(2026, 7, 20),
            due_date=date(2026, 7, 25),
            scheduled_date=date(2026, 7, 25),
            payment_method="Boletos de Outros Bancos",
            document_type="Nota Fiscal",
            responsible="FRANCISCO LUAN",
            attachments="NF 1001.pdf | BOLETO 1001.pdf",
            attachment_count=2,
        )

        doc2_no_boleto = CmflexOpenDocument(
            document_id=2,
            company_id=1,
            company="CARMEL TAÍBA",
            launch_number="2002",
            installment="1",
            supplier="MANDACARU SERVICOS LTDA",
            tax_id="09370323000141",
            reference="NF 2002",
            amount=Decimal("2000.00"),
            balance=Decimal("2000.00"),
            launch_date=date(2026, 7, 20),
            due_date=date(2026, 7, 26),
            scheduled_date=date(2026, 7, 26),
            payment_method="Transferência Bancária",
            document_type="Nota Fiscal",
            responsible="DEVID SILVA",
            attachments="NF 2002.pdf",
            attachment_count=1,
        )

        items = reconcile_dda((self.tx1, self.tx2, self.tx_unmatched), (doc1, doc2_no_boleto))
        self.assertEqual(len(items), 3)

        # tx1 -> Matched with doc1, has boleto -> LANÇ and SEM ALTERAÇÃO
        self.assertEqual(items[0].effective_status, STATUS_LANC)
        self.assertEqual(items[0].effective_observation, OBS_SEM_ALTERACAO)
        self.assertEqual(items[0].effective_responsible, "FRANCISCO LUAN")
        self.assertEqual(items[0].matched_document, doc1)

        # tx2 -> Matched with doc2, transfer method without boleto -> SEM BOLETO E DADOS BAN.
        self.assertEqual(items[1].effective_status, STATUS_LANC)
        self.assertEqual(items[1].effective_observation, OBS_SEM_BOLETO_E_DADOS_BAN)
        self.assertEqual(items[1].effective_responsible, "DEVID SILVA")

        # tx_unmatched -> Not in ERP -> NÃO LANÇ.
        self.assertEqual(items[2].effective_status, STATUS_NAO_LANC)
        self.assertEqual(items[2].effective_observation, OBS_NAO_LOCALIZADO)
        self.assertIsNone(items[2].matched_document)

    def test_reconciliation_detects_oc(self):
        doc_oc = CmflexOpenDocument(
            document_id=3,
            company_id=1,
            company="CARMEL TAÍBA",
            launch_number="1001",
            installment="1",
            supplier="FORNECEDOR ALPHA COMERCIO LTDA",
            tax_id="07647045000150",
            reference="",
            amount=Decimal("1500.00"),
            balance=Decimal("1500.00"),
            launch_date=date(2026, 7, 20),
            due_date=date(2026, 7, 25),
            scheduled_date=date(2026, 7, 25),
            payment_method="Boletos de Próprio Banco",
            document_type="Ordem de Compra",
            responsible="DOUGLAS TORQUATO",
            attachments="BOLETO.pdf",
            attachment_count=1,
        )
        items = reconcile_dda((self.tx1,), (doc_oc,))
        self.assertEqual(items[0].effective_status, STATUS_LANC_OC)


class DdaStoreTests(TestCase):
    def test_store_and_apply_reviews(self):
        with TemporaryDirectory() as tmp_dir:
            db_path = Path(tmp_dir) / "test_dda.sqlite3"
            store = DdaStore(db_path)

            tx = DdaTransaction(
                sheet_name="CHARME",
                row_index=5,
                company="CARMEL CHARME",
                bank_code="341",
                tax_id="12.345.678/0001-90",
                supplier="FORNECEDOR TESTE",
                document_number="5005",
                charge_type="Cobrança",
                situation="A pagar",
                issue_date=date(2026, 7, 20),
                due_date=date(2026, 7, 25),
                amount=Decimal("500.00"),
            )

            store.save_review(
                tx,
                custom_status=STATUS_LANC,
                custom_observation="Ajustado manualmente",
                custom_responsible="VICTOR MIRANDA",
            )

            reviews = store.load_reviews()
            key = make_transaction_key(tx)
            self.assertIn(key, reviews)
            self.assertEqual(reviews[key], (STATUS_LANC, "Ajustado manualmente", "VICTOR MIRANDA"))

            # Test apply
            item = DdaReconciliationItem(transaction=tx, status=STATUS_NAO_LANC)
            store.apply_reviews([item])
            self.assertEqual(item.effective_status, STATUS_LANC)
            self.assertEqual(item.effective_observation, "Ajustado manualmente")
            self.assertEqual(item.effective_responsible, "VICTOR MIRANDA")


class DdaExportTests(TestCase):
    def test_excel_export_and_email_generation(self):
        tx1 = DdaTransaction(
            sheet_name="TAIBA",
            row_index=2,
            company="CARMEL TAÍBA",
            bank_code="341",
            tax_id="07.647.045/0001-50",
            supplier="FORNECEDOR ALPHA",
            document_number="1001",
            charge_type="Cobrança",
            situation="A pagar",
            issue_date=date(2026, 7, 20),
            due_date=date(2026, 7, 25),
            amount=Decimal("1500.00"),
        )
        item1 = DdaReconciliationItem(
            transaction=tx1,
            status=STATUS_NAO_LANC,
            observation=OBS_NAO_LOCALIZADO,
            responsible="FRANCISCO LUAN",
        )
        tx2 = DdaTransaction(
            sheet_name="CHARME",
            row_index=3,
            company="CARMEL CHARME",
            bank_code="237",
            tax_id="09.370.323/0001-41",
            supplier="MANDACARU",
            document_number="2002",
            charge_type="Cobrança",
            situation="A pagar",
            issue_date=date(2026, 7, 20),
            due_date=date(2026, 7, 26),
            amount=Decimal("2000.00"),
        )
        item2 = DdaReconciliationItem(
            transaction=tx2,
            status=STATUS_LANC,
            observation=OBS_SEM_BOLETO_E_DADOS_BAN,
            responsible="DEVID SILVA",
        )

        excel_data = export_dda_reconciliation((item1, item2))
        self.assertIsInstance(excel_data, bytes)

        wb = load_workbook(BytesIO(excel_data))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "EMPRESA")
        self.assertEqual(ws.cell(row=1, column=13).value, "STATUS")
        self.assertEqual(ws.cell(row=1, column=14).value, "OBSERVAÇÃO")
        self.assertEqual(ws.cell(row=1, column=15).value, "RESPONSAVEL")

        self.assertEqual(ws.cell(row=2, column=13).value, STATUS_NAO_LANC)
        self.assertEqual(ws.cell(row=3, column=13).value, STATUS_LANC)

        # Email generator
        email_text = generate_collection_email((item1, item2))
        self.assertIn("FRANCISCO LUAN", email_text)
        self.assertIn("DEVID SILVA", email_text)
        self.assertIn("1.500,00", email_text)
        self.assertIn("2.000,00", email_text)


class DdaPanelTests(TestCase):
    def _walk(self, control: ft.Control):
        yield control
        for attr in ("content", "controls"):
            child = getattr(control, attr, None)
            if isinstance(child, ft.Control):
                yield from self._walk(child)
            elif isinstance(child, (list, tuple)):
                for item in child:
                    if isinstance(item, ft.Control):
                        yield from self._walk(item)

    def test_panel_initial_state_and_upload(self):
        page = MagicMock(web=True)
        page.overlay = []
        page.get_upload_url.return_value = "/upload?token=test"
        page.run_task.side_effect = lambda handler, *args: asyncio.run(handler(*args))

        with TemporaryDirectory() as tmp_dir:
            store = MagicMock()
            store.load_reviews.return_value = {}
            provider = MagicMock()
            provider.fetch_open_documents.return_value = ()

            panel_view = build_dda_panel(
                page,
                download_file=MagicMock(),
                provider=provider,
                upload_dir=tmp_dir,
                store=store,
            )

            # Initially shows empty/upload screen
            self.assertIsNotNone(panel_view.control)

            # Simulate upload of a DDA workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "TAIBA"
            ws.append([
                "EMPRESA", "BANCO", "CNPJ", "FORNECEDOR", "DOCUMENTO", "TIPO", "SITUAÇÃO", "DATA EMISSÃO", "DATA VENCIMENTO", "VALOR"
            ])
            ws.append([
                "CARMEL TAÍBA", "341", "07.647.045/0001-50", "FORNECEDOR ALPHA", "1001", "Cobrança", "A pagar", "20/07/2026", "25/07/2026", "1.500,00"
            ])
            file_path = Path(tmp_dir) / "dda_test.xlsx"
            wb.save(file_path)

            picker = page.overlay[0]
            picker.upload = MagicMock()
            picker.on_result(MagicMock(files=[SimpleNamespace(name="dda_test.xlsx", size=5000)]))

            # Trigger upload progress complete
            picker.on_upload(MagicMock(file_name="dda_test.xlsx", progress=1.0, error=None))

            # Database should have been queried
            provider.fetch_open_documents.assert_called_once()

            # DataTable should now exist
            table = next(
                item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable)
            )
            self.assertEqual(len(table.rows), 1)

            # Test filter interactions
            dropdowns = [item for item in self._walk(panel_view.control) if isinstance(item, ft.Dropdown)]
            status_dd = next(dd for dd in dropdowns if dd.label == "Situação")
            # Change status filter to LANÇ (should produce 0 rows since tx is NÃO LANÇ.)
            status_dd.on_change(SimpleNamespace(control=SimpleNamespace(value=STATUS_LANC)))
            table_after = next(item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table_after.rows), 0)

            # Change status filter back to Todos
            status_dd.on_change(SimpleNamespace(control=SimpleNamespace(value="Todos")))
            table_after2 = next(item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table_after2.rows), 1)

            # Test search filter
            text_fields = [item for item in self._walk(panel_view.control) if isinstance(item, ft.TextField)]
            search_tf = next(tf for tf in text_fields if "Buscar" in (tf.label or ""))
            search_tf.on_change(SimpleNamespace(control=SimpleNamespace(value="ALPHA")))
            table_after3 = next(item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table_after3.rows), 1)

            search_tf.on_change(SimpleNamespace(control=SimpleNamespace(value="INEXISTENTE")))
            table_after4 = next(item for item in self._walk(panel_view.control) if isinstance(item, ft.DataTable))
            self.assertEqual(len(table_after4.rows), 0)
