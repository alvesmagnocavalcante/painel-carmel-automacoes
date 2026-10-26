"""Leitura da fatura e controles de conciliação sem acesso ao Itaú."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import MagicMock, patch

import flet as ft
from openpyxl import load_workbook

from automations.cartoes_credito.bank_panel import main as build_bank_panel
from automations.cartoes_credito.export import export_statement
from automations.cartoes_credito.reconciliation import (
    CARD_CHARGE_STATUS, BankReview, classify_transaction, filter_transactions,
    is_card_charge, suggest_matches,
)
from automations.cartoes_credito.statement import BankStatement, BankTransaction, parse_rows, parse_statement
from automations.cartoes_credito.store import CardStore
from tests.test_cartoes import payment, walk
from portal.theme import DARK


def statement_rows():
    rows = [[None] * 11 for _ in range(30)]
    rows[2][0] = "MAGNA PRAIA HOTEL LTDA"
    rows[5][0] = "ITAU EMPRESAS VISA - cartão mascarado"
    rows[8][2] = date(2026, 10, 11)
    rows[12][10], rows[13][10], rows[14][10], rows[15][10] = 90, 50, 0, 140
    rows[20][0] = "TITULAR - FINAL 407505******1528"
    rows[22][0] = "Lançamentos nacionais"
    rows[23][0], rows[23][2], rows[23][10] = date(2026, 9, 1), "COMPRA", 100
    rows[24][0], rows[24][2], rows[24][10] = date(2026, 9, 2), "ESTORNO", -10
    rows[25][0] = "Total de lançamentos nacionais"
    rows[26][0] = "Lançamentos internacionais"
    rows[27][0], rows[27][2], rows[27][10] = date(2026, 9, 3), "COMPRA EXTERIOR", 50
    rows[28][0] = "Total de lançamentos internacionais"
    return [tuple(row) for row in rows]


class StatementTests(TestCase):
    def test_parser_keeps_transactions_and_reconciles_summary(self):
        result = parse_rows(statement_rows())
        self.assertEqual(len(result.transactions), 3)
        self.assertEqual(result.total, Decimal("140.00"))
        self.assertEqual(result.transactions[0].card_last_four, "1528")
        self.assertEqual(result.transactions[1].amount, Decimal("-10.00"))
        self.assertEqual(result.transactions[2].category, "Internacional")

    def test_parser_rejects_incomplete_extraction_and_bad_file(self):
        rows = statement_rows()
        rows[12] = (*rows[12][:10], 91)
        with self.assertRaisesRegex(ValueError, "subtotais"):
            parse_rows(rows)
        with self.assertRaisesRegex(ValueError, "inválida"):
            parse_statement(b"not an xlsx")

    def test_transaction_filters_only_change_visible_rows(self):
        statement = parse_rows(statement_rows())
        fee = BankTransaction(30, date(2026, 9, 4), "ANUIDADE", Decimal("20.50"),
                              "TITULAR", "1528", "Nacional")
        transactions = (*statement.transactions, fee)
        matched = payment(company="MAGNA PRAIA", amount=Decimal("100"))
        other_company = payment(company="CARMEL TAÍBA", amount=Decimal("50"))
        rows = (matched, other_company)
        self.assertEqual(filter_transactions(transactions, rows, "MAGNA PRAIA", "Todos"), transactions)
        self.assertEqual(filter_transactions(transactions, rows, "MAGNA PRAIA",
                                             "Com lançamento na busca"), (transactions[0],))
        self.assertEqual(filter_transactions(transactions, rows, "MAGNA PRAIA",
                                             "Sem lançamento na busca"), transactions[1:])
        self.assertEqual(filter_transactions(transactions, rows, "MAGNA PRAIA", "Encargos"), (fee,))
        self.assertEqual(filter_transactions(transactions, rows, "Todos", "Sem lançamento na busca"), ())

    def test_suggestions_are_only_candidates_and_confirmation_is_guarded(self):
        transaction = BankTransaction(24, date(2026, 9, 1), "COMPRA", Decimal("100"), "TITULAR", "1528", "Nacional")
        good = payment(document_id=10, company="MAGNA PRAIA", amount=Decimal("100"), launch_date=date(2026, 9, 20))
        late = payment(document_id=11, company="MAGNA PRAIA", amount=Decimal("100"), launch_date=date(2026, 9, 1) + timedelta(days=61))
        wrong = payment(document_id=12, company="MAGNA PRAIA", amount=Decimal("99"))
        rows = (good, late, wrong)
        self.assertEqual(suggest_matches(transaction, rows, "MAGNA PRAIA"), (good, late))
        self.assertEqual(suggest_matches(transaction, rows, "Todos"), ())
        with self.assertRaisesRegex(ValueError, "Associe"):
            classify_transaction(transaction, "Conferido", "", None, rows, "MAGNA PRAIA", {})
        with self.assertRaisesRegex(ValueError, "diverge"):
            classify_transaction(transaction, "Conferido", "", wrong.key, rows, "MAGNA PRAIA", {})
        review = classify_transaction(transaction, "Conferido", "", good.key, rows, "MAGNA PRAIA", {})
        self.assertEqual(review, BankReview("Conferido", "", good.key))
        with self.assertRaisesRegex(ValueError, "já foi confirmado"):
            classify_transaction(BankTransaction(25, transaction.purchase_date, "OUTRA", transaction.amount,
                                                "TITULAR", "1528", "Nacional"),
                                 "Conferido", "", good.key, rows, "MAGNA PRAIA", {24: review})

    def test_card_charge_is_confirmed_without_internal_payment_only_after_review(self):
        fee = BankTransaction(24, date(2026, 8, 30), "ANUIDADE DIFERENCI09/12",
                              Decimal("20.50"), "TITULAR", "1528", "Nacional")
        normal = BankTransaction(25, fee.purchase_date, "COMPRA", fee.amount,
                                 fee.cardholder, fee.card_last_four, fee.category)
        self.assertTrue(is_card_charge(fee))
        self.assertTrue(is_card_charge(BankTransaction(26, fee.purchase_date, "IOF COMPRA INTERNACIONAL",
                                                     fee.amount, fee.cardholder, fee.card_last_four, fee.category)))
        self.assertTrue(is_card_charge(BankTransaction(27, fee.purchase_date, "JUROS DO CARTÃO",
                                                     fee.amount, fee.cardholder, fee.card_last_four, "Encargo")))
        self.assertFalse(is_card_charge(normal))
        self.assertEqual(classify_transaction(fee, CARD_CHARGE_STATUS, "", None, (), "MAGNA PRAIA", {}),
                         BankReview(CARD_CHARGE_STATUS))
        with self.assertRaisesRegex(ValueError, "Apenas anuidades"):
            classify_transaction(normal, CARD_CHARGE_STATUS, "", None, (), "MAGNA PRAIA", {})
        linked = payment(amount=fee.amount, company="MAGNA PRAIA")
        with self.assertRaisesRegex(ValueError, "Desmarque"):
            classify_transaction(fee, CARD_CHARGE_STATUS, "", linked.key, (linked,), "MAGNA PRAIA", {})

        statement = BankStatement("MAGNA PRAIA", date(2026, 10, 11), fee.amount, (fee,))
        sheet = load_workbook(BytesIO(export_statement(statement, {24: BankReview(CARD_CHARGE_STATUS)}, ()))).active
        self.assertEqual(sheet["N2"].value, CARD_CHARGE_STATUS)
        self.assertIsNone(sheet["I2"].value)
        self.assertEqual(sheet["G3"].value, "=SUM(G2:G2)")
        self.assertEqual(sheet["G4"].value, 20.5)

    def test_export_retains_source_rows_and_manual_outcomes(self):
        statement = parse_rows(statement_rows())
        matched = payment(document_id=10, company="MAGNA PRAIA", amount=Decimal("100"))
        content = export_statement(statement, {24: BankReview("Conferido", "", matched.key)}, (matched,))
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet["A2"].value, 24)
        self.assertEqual(sheet["N2"].value, "Conferido")
        self.assertEqual(sheet["I2"].value, matched.launch_number)
        self.assertEqual(sheet["G5"].value, "=SUM(G2:G4)")
        self.assertEqual(sheet["G6"].value, 140)

    def test_upload_is_consumed_and_not_kept_on_disk(self):
        statement = parse_rows(statement_rows())
        page = MagicMock(web=True, width=760)
        page.overlay = []
        page.get_upload_url.return_value = "/upload?signature=test"
        on_loaded = MagicMock()
        with TemporaryDirectory() as folder, patch("automations.cartoes_credito.bank_panel.parse_statement", return_value=statement):
            view = build_bank_panel(page, colors=lambda: DARK,
                                    payments=lambda: (payment(company="MAGNA PRAIA", amount=Decimal("99")),),
                                    company=lambda: "MAGNA PRAIA",
                                    lookup_state=lambda: "", on_loaded=on_loaded,
                                    download_file=MagicMock(), upload_dir=folder,
                                    store=CardStore(Path(folder) / "cartoes.sqlite3"))
            picker = page.overlay[0]
            picker.upload = MagicMock()
            picked = MagicMock(files=[MagicMock(name="fatura.xlsx", size=1234)])
            picker.on_result(picked)
            token = page.get_upload_url.call_args.args[0]
            path = Path(folder) / token
            path.write_bytes(b"temporary workbook")
            picker.on_upload(MagicMock(error=None, progress=1.0))
            self.assertIs(view.statement(), statement)
            self.assertFalse(path.exists())
            on_loaded.assert_called_once_with(statement)
            rows = view.control.content.controls[-1]
            self.assertIsInstance(rows, ft.Column)
            self.assertIsNone(rows.scroll)
            self.assertEqual(view.control.content.horizontal_alignment, ft.CrossAxisAlignment.STRETCH)
            self.assertEqual(len(rows.controls), len(statement.transactions) + 1)
            filter_field = next(item for item in walk(view.control)
                                if isinstance(item, ft.Dropdown) and item.label == "Filtrar itens")
            filter_field.value = "Sem lançamento na busca"
            filter_field.on_change(None)
            self.assertEqual(len(view.control.content.controls[-1].controls), len(statement.transactions) + 1)
            filter_field.value = "Com lançamento na busca"
            filter_field.on_change(None)
            self.assertEqual(len(view.control.content.controls[-1].controls), 1)
            filter_field.value = "Todos"
            filter_field.on_change(None)
            self.assertFalse(rows.controls[1].content.wrap)
            action = next(item for item in walk(rows.controls[1]) if isinstance(item, ft.OutlinedButton))
            action.on_click(None)
            dialog = page.open.call_args.args[0]
            self.assertEqual(dialog.content.width, 712)
            self.assertEqual(dialog.content.height, 370)
            self.assertFalse(any(isinstance(item, ft.Text) and "CMFlex" in item.value
                                 for item in walk(dialog)))
            self.assertTrue(any(isinstance(item, ft.Text) and
                                "Esta cobrança consta na fatura" in item.value
                                for item in walk(dialog)))
            self.assertFalse(any(isinstance(item, ft.RadioGroup) for item in dialog.content.content.controls))
            result = next(item for item in walk(dialog) if isinstance(item, ft.Dropdown)
                          and item.label == "Resultado")
            self.assertEqual([option.key for option in result.options], ["A conferir", "Pendente"])
            result.value = "Pendente"
            result.on_change(None)
            reason = next(item for item in walk(dialog) if isinstance(item, ft.TextField)
                          and item.label == "Pendência")
            reason.value = "Estorno sem lançamento correspondente"
            dialog.actions[-1].on_click(None)
            export = next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton)
                          and item.text == "Baixar parcial")
            self.assertFalse(export.disabled)
            self.assertTrue(dialog._build_add_commands())
            self.assertTrue(view.control._build_add_commands())

    def test_card_charge_has_a_clear_review_path_in_the_dialog(self):
        fee = BankTransaction(24, date(2026, 8, 30), "ANUIDADE DIFERENCI09/12",
                              Decimal("20.50"), "TITULAR", "1528", "Nacional")
        statement = BankStatement("MAGNA PRAIA", date(2026, 10, 11), fee.amount, (fee,))
        page = MagicMock(web=True, width=1200)
        page.overlay = []
        page.get_upload_url.return_value = "/upload?signature=test"
        with TemporaryDirectory() as folder, patch("automations.cartoes_credito.bank_panel.parse_statement",
                                                   return_value=statement):
            view = build_bank_panel(page, colors=lambda: DARK, payments=lambda: (),
                                    company=lambda: "MAGNA PRAIA", lookup_state=lambda: "",
                                    on_loaded=lambda _: None, download_file=MagicMock(), upload_dir=folder,
                                    store=CardStore(Path(folder) / "cartoes.sqlite3"))
            picker = page.overlay[0]
            picker.upload = MagicMock()
            picker.on_result(MagicMock(files=[MagicMock(name="fatura.xlsx", size=1234)]))
            path = Path(folder) / page.get_upload_url.call_args.args[0]
            path.write_bytes(b"temporary workbook")
            picker.on_upload(MagicMock(error=None, progress=1.0))

            action = next(item for item in walk(view.control) if isinstance(item, ft.OutlinedButton)
                          and item.text == "Conferir")
            action.on_click(None)
            dialog = page.open.call_args.args[0]
            self.assertEqual(dialog.title.value, "Conferir item da fatura")
            self.assertEqual(dialog.content.height, 430)
            self.assertTrue(any(isinstance(item, ft.Text) and item.value == "Encargo do cartão"
                                for item in walk(dialog)))
            self.assertTrue(any(isinstance(item, ft.Text) and
                                "Não encontramos lançamento interno de R$ 20,50" in item.value
                                for item in walk(dialog)))
            result = next(item for item in walk(dialog) if isinstance(item, ft.Dropdown)
                          and item.label == "Resultado")
            self.assertEqual([option.key for option in result.options],
                             ["A conferir", CARD_CHARGE_STATUS, "Pendente"])
            self.assertTrue(dialog._build_add_commands())
            result.value = CARD_CHARGE_STATUS
            result.on_change(None)
            dialog.actions[-1].on_click(None)
            self.assertTrue(any(isinstance(item, ft.Text) and item.value == CARD_CHARGE_STATUS
                                for item in walk(view.control)))
            self.assertTrue(any(isinstance(item, ft.Text) and "1 de 1 itens analisados" in item.value
                                for item in walk(view.control)))
