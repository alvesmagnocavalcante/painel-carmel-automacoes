"""Data and export services used by the Bordero panel."""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal
from pathlib import Path
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def format_tax_id(value: object) -> str:
    """Format numeric CPF/CNPJ without padding, guessing, or validating check digits."""
    text = "" if value is None else str(value).strip()
    if not text:
        return "Não informado"
    if not re.fullmatch(r"[0-9.\-/\s]+", text):
        return text
    digits = re.sub(r"[^0-9]", "", text)
    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    if len(digits) == 14:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"
    return text


@dataclass(frozen=True)
class ConferenceRow:
    supplier: str
    tax_id: str
    launch_number: str
    company: str
    balance: Decimal
    scheduled_date: date | None
    payment_method: str
    history: str
    situation: str
    responsible: str
    due_date: date | None = None
    details: str = ""
    review_key: str = ""
    pending_reason: str = ""
    attachment_count: int | None = None
    source_situation: str = ""
    attachment_names: str = ""
    document_id: int | None = None
    company_id: int | None = None

    @property
    def attachment_summary(self) -> str:
        if self.attachment_count is None:
            return "Não informado"
        return f"Sim ({self.attachment_count})" if self.attachment_count > 0 else "Sem anexo"


def classify_conference(row: ConferenceRow, status: str, reason: str = "") -> ConferenceRow:
    if status not in {"Conferido", "Pendente", "A conferir"}:
        raise ValueError("Classificação inválida.")
    if status == "Pendente" and not reason.strip():
        raise ValueError("Informe o motivo da pendência.")
    return replace(row, situation=status, pending_reason=reason.strip() if status == "Pendente" else "")


def filter_conferences(
    rows: tuple[ConferenceRow, ...], classification: str = "Todos",
    payment_method: str | None = None, search: str = "",
) -> tuple[ConferenceRow, ...]:
    def normalize(value: str) -> str:
        return " ".join("".join(character for character in unicodedata.normalize("NFKD", value.casefold())
                                if not unicodedata.combining(character)).split())

    query = normalize(search)
    # Numeric document searches accept masks, but letters must not be discarded.
    digits = re.sub(r"[^0-9]", "", query) if re.fullmatch(r"[0-9.\-/\s]+", query) else ""
    return tuple(
        row for row in rows
        if (classification == "Todos" or row.situation == classification)
        and (payment_method is None or row.payment_method == payment_method)
        and (not query or query in normalize(row.supplier) or query in normalize(row.launch_number)
             or query in normalize(row.tax_id)
             or bool(digits and digits in re.sub(r"[^0-9]", "", row.tax_id)))
    )


def validate_filters(start_date: date, end_date: date, hotel: str) -> None:
    if start_date > end_date:
        raise ValueError("A data inicial n\u00e3o pode ser posterior \u00e0 data final.")
    if not hotel.strip():
        raise ValueError("Selecione uma empresa ou Todos.")
    if (end_date - start_date).days > 366:
        raise ValueError("O per\u00edodo m\u00e1ximo permitido \u00e9 de 366 dias.")


def export_conferences(
    rows: tuple[ConferenceRow, ...],
    destination: Path,
    *,
    hotel: str,
    start_date: date,
    end_date: date,
) -> Path:
    if not rows:
        raise ValueError("N\u00e3o existem confer\u00eancias para exportar.")

    destination = destination.expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Confer\u00eancias"

    worksheet.append(["Confer\u00eancia de pagamentos em aberto"])
    worksheet.append(["Hotel", hotel])
    worksheet.append(
        ["Per\u00edodo", f"{start_date:%d/%m/%Y} a {end_date:%d/%m/%Y}"]
    )
    worksheet.append([])
    headers = (
        "Fornecedor",
        "CPF/CNPJ",
        "N\u00famero do lan\u00e7amento",
        "Empresa",
        "Saldo",
        "Data programada",
        "Forma de pagamento",
        "Hist\u00f3rico",
        "Situa\u00e7\u00e3o",
        "Respons\u00e1vel",
        "Vencimento",
        "Motivo da pend\u00eancia",
        "Anexos",
        "Situação BPM",
    )
    worksheet.append(headers)

    for row in rows:
        worksheet.append(
            [
                row.supplier,
                row.tax_id,
                row.launch_number,
                row.company,
                float(row.balance),
                row.scheduled_date,
                row.payment_method,
                row.history,
                row.situation,
                row.responsible,
                row.due_date,
                row.pending_reason,
                row.attachment_summary,
                row.source_situation,
            ]
        )

    # Database and review text must remain literal, never Excel formulas/errors.
    for cells in worksheet.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    title = worksheet["A1"]
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="183B56")
    worksheet.merge_cells("A1:N1")
    worksheet["A1"].alignment = Alignment(horizontal="left")

    for cell in worksheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6F5F")
        cell.alignment = Alignment(horizontal="center")

    for row_number in range(6, worksheet.max_row + 1):
        worksheet.cell(row_number, 5).number_format = 'R$ #,##0.00'
        worksheet.cell(row_number, 6).number_format = "dd/mm/yyyy"
        worksheet.cell(row_number, 11).number_format = "dd/mm/yyyy"

    widths = (30, 21, 22, 22, 16, 18, 25, 42, 25, 22, 18, 45, 18, 25)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = f"A5:N{worksheet.max_row}"

    temporary = destination.with_name(f".{destination.name}.tmp")
    try:
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return destination
