"""Data structures and parser for Flash corporate card statements."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Sequence

MONTH_NAMES = (
    "JANEIRO",
    "FEVEREIRO",
    "MARÇO",
    "ABRIL",
    "MAIO",
    "JUNHO",
    "JULHO",
    "AGOSTO",
    "SETEMBRO",
    "OUTUBRO",
    "NOVEMBRO",
    "DEZEMBRO",
)

STATUS_PENDING_LAUNCH = "PENDENTE DE LANÇAMENTO"
STATUS_PENDING_SETTLEMENT = "PENDENTE DE PRESTAÇÃO"
STATUS_SETTLED = "PRESTADO CONTAS"


def normalize_text(value: str | None) -> str:
    """Normalize text for insensitive comparison removing accents and double spaces."""
    if not value:
        return ""
    normalized = unicodedata.normalize("NFKD", str(value).strip().casefold())
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return " ".join(without_marks.split())


def get_month_name(target_date: date) -> str:
    """Return the uppercase Portuguese month name."""
    return MONTH_NAMES[target_date.month - 1]


def parse_flash_date(value: object) -> date | None:
    """Extract a date object from text, date, or datetime object."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    raw = str(value).strip()
    if not raw:
        return None

    # Standard date formats
    for fmt in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass

    # Regex search for DD/MM/YYYY
    match_br = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if match_br:
        day, month, year = match_br.groups()
        try:
            return date(int(year), int(month), int(day))
        except ValueError:
            pass

    # Regex search for YYYY-MM-DD
    match_iso = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", raw)
    if match_iso:
        year, month, day = match_iso.groups()
        try:
            return date(int(year), int(month), int(day))
        except ValueError:
            pass

    return None


def parse_flash_amount(value: object) -> tuple[Decimal | None, bool]:
    """Parse monetary value returning (positive_amount, is_credit_or_deposit).

    Example inputs:
      "- R$ 100,00" -> (Decimal('100.00'), False)
      "+ R$ 90,90"  -> (Decimal('90.90'), True)
      "-85.4"       -> (Decimal('85.40'), False)
    """
    if value is None:
        return None, False
    if isinstance(value, (int, float, Decimal)):
        dec = Decimal(str(value))
        is_credit = dec > 0
        return abs(dec).quantize(Decimal("0.01")), is_credit

    raw = str(value).strip().replace("\xa0", " ")
    if not raw:
        return None, False

    is_credit = "+" in raw or (not raw.startswith("-") and "deposito" in normalize_text(raw))
    cleaned = re.sub(r"[^0-9,.-]", "", raw)
    if not cleaned:
        return None, False

    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")

    try:
        dec = Decimal(cleaned)
        is_negative = dec < 0 or raw.startswith("-")
        return abs(dec).quantize(Decimal("0.01")), not is_negative if (dec != 0) else is_credit
    except InvalidOperation:
        return None, False


@dataclass(frozen=True)
class FlashTransaction:
    """A transaction extracted from the Flash card statement."""

    row_index: int
    purchase_date: date
    raw_date: str
    description: str
    amount: Decimal
    person: str
    payment_method: str = ""
    statement_status: str = ""
    is_deposit: bool = False
    hotel: str = ""

    @property
    def month_name(self) -> str:
        return get_month_name(self.purchase_date)


@dataclass(frozen=True)
class CmflexFlashDocument:
    """A document registered in CMFlex under cmflex.historico_cartao_flash."""

    document_id: int
    company_id: int
    company: str
    launch_number: str
    supplier: str
    tax_id: str
    launch_date: date | None
    due_date: date | None
    scheduled_date: date | None
    amount: Decimal
    balance: Decimal
    document_status: int
    payment_method: str
    holder: str
    history: str = ""
    attachment_count: int = 0

    @property
    def is_settled(self) -> bool:
        """Document is settled if balance is zero or status is 2 (baixado)."""
        return self.balance == Decimal("0.0") or self.document_status == 2


@dataclass(frozen=True)
class FlashReconciliationItem:
    """An item in the reconciliation between Flash and CMFlex."""

    transaction: FlashTransaction
    status: str
    matched_document: CmflexFlashDocument | None = None
    manual_notes: str = ""
    hotel: str = ""

    @property
    def effective_hotel(self) -> str:
        if self.hotel:
            return self.hotel
        if self.matched_document and self.matched_document.company:
            return self.matched_document.company
        return self.transaction.hotel or "Não informada"

    @property
    def effective_notes(self) -> str:
        if self.manual_notes:
            return self.manual_notes
        return self.status


@dataclass(frozen=True)
class FlashStatement:
    """Collection of transactions parsed from a Flash statement file."""

    transactions: tuple[FlashTransaction, ...]
    file_name: str = ""
    total_expenses: Decimal = field(init=False)
    total_deposits: Decimal = field(init=False)

    def __post_init__(self) -> None:
        expenses = sum((t.amount for t in self.transactions if not t.is_deposit), Decimal("0.00"))
        deposits = sum((t.amount for t in self.transactions if t.is_deposit), Decimal("0.00"))
        object.__setattr__(self, "total_expenses", expenses.quantize(Decimal("0.01")))
        object.__setattr__(self, "total_deposits", deposits.quantize(Decimal("0.01")))

    @property
    def expense_transactions(self) -> tuple[FlashTransaction, ...]:
        return tuple(t for t in self.transactions if not t.is_deposit)


def _map_flash_columns(header: Sequence[str]) -> dict[str, int]:
    """Map column indices by recognizing standard and alternate Flash column names."""
    indices: dict[str, int] = {}
    for i, col in enumerate(header):
        normalized = normalize_text(col)
        if any(k in normalized for k in ("data", "date", "dt")) and "date" not in indices:
            indices["date"] = i
        elif any(k in normalized for k in ("movimentacao", "descricao", "estabelecimento", "fornecedor", "historico", "transacao")):
            indices["desc"] = i
        elif any(k in normalized for k in ("valor", "vl", "amount", "total")):
            indices["amount"] = i
        elif any(k in normalized for k in ("pessoa", "portador", "colaborador", "nome", "titular", "usuario")):
            indices["person"] = i
        elif any(k in normalized for k in ("pagamento", "forma", "cartao", "meio")):
            indices["payment"] = i
        elif any(k in normalized for k in ("prestacao", "status", "situacao")):
            indices["status"] = i
        elif any(k in normalized for k in ("hotel", "empresa", "unidade", "filial")):
            indices["hotel"] = i
    return indices


def _detect_columns_by_content(sample_rows: Sequence[Sequence[str]]) -> dict[str, int]:
    """Detect columns by inspecting sample row values when no header row is present."""
    detected: dict[str, int] = {}
    if not sample_rows:
        return detected

    col_count = max(len(r) for r in sample_rows)
    for col_idx in range(col_count):
        date_hits = 0
        amount_hits = 0
        for r in sample_rows:
            if col_idx >= len(r):
                continue
            val = r[col_idx]
            if parse_flash_date(val) is not None:
                date_hits += 1
            amt, _ = parse_flash_amount(val)
            if amt is not None:
                amount_hits += 1

        if date_hits >= max(1, len(sample_rows) // 2) and "date" not in detected:
            detected["date"] = col_idx
        elif amount_hits >= max(1, len(sample_rows) // 2) and "amount" not in detected:
            detected["amount"] = col_idx

    # If we found date and amount, find text columns for desc and person
    remaining = [i for i in range(col_count) if i != detected.get("date") and i != detected.get("amount")]
    for i in remaining:
        # Check if first cell is a month name
        first_val = normalize_text(sample_rows[0][i]) if sample_rows and i < len(sample_rows[0]) else ""
        if any(m.lower() in first_val for m in MONTH_NAMES):
            continue
        if "desc" not in detected:
            detected["desc"] = i
        elif "person" not in detected:
            detected["person"] = i
        elif "hotel" not in detected:
            detected["hotel"] = i

    return detected


def _build_transaction_from_row(
    row: Sequence[object], indices: dict[str, int], row_index: int
) -> FlashTransaction | None:
    if "date" not in indices or "amount" not in indices:
        return None

    date_idx = indices["date"]
    date_val = row[date_idx] if date_idx < len(row) else ""
    parsed_date = parse_flash_date(date_val)
    if not parsed_date:
        return None

    amount_idx = indices["amount"]
    amount_str = row[amount_idx] if amount_idx < len(row) else ""
    amount, is_credit = parse_flash_amount(amount_str)
    if amount is None or amount == Decimal("0.00"):
        return None

    desc_idx = indices.get("desc")
    desc = str(row[desc_idx]).strip() if (desc_idx is not None and desc_idx < len(row) and row[desc_idx] is not None) else "Despesa Cartão Flash"

    person_idx = indices.get("person")
    person = str(row[person_idx]).strip() if (person_idx is not None and person_idx < len(row) and row[person_idx] is not None) else ""

    payment_idx = indices.get("payment")
    payment = str(row[payment_idx]).strip() if (payment_idx is not None and payment_idx < len(row) and row[payment_idx] is not None) else ""

    status_idx = indices.get("status")
    status = str(row[status_idx]).strip() if (status_idx is not None and status_idx < len(row) and row[status_idx] is not None) else ""

    hotel_idx = indices.get("hotel")
    hotel = str(row[hotel_idx]).strip() if (hotel_idx is not None and hotel_idx < len(row) and row[hotel_idx] is not None) else ""

    is_deposit = is_credit or "deposito" in normalize_text(desc)

    return FlashTransaction(
        row_index=row_index,
        purchase_date=parsed_date,
        raw_date=str(date_val),
        description=desc,
        amount=amount,
        person=person,
        payment_method=payment,
        statement_status=status,
        is_deposit=is_deposit,
        hotel=hotel,
    )


def parse_flash_csv(content: str | bytes, file_name: str = "extrato.csv") -> FlashStatement:
    """Parse CSV statement exported from the Flash web portal."""
    if isinstance(content, bytes):
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = content.decode("utf-8", errors="replace")
    else:
        text = content

    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return FlashStatement(transactions=(), file_name=file_name)

    sample = "\n".join(lines[:10])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ";" if ";" in lines[0] else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    all_rows = [list(r) for r in reader if any(r)]
    if not all_rows:
        return FlashStatement(transactions=(), file_name=file_name)

    # Search for header row
    header_idx = -1
    col_indices: dict[str, int] = {}
    for idx, row in enumerate(all_rows[:15]):
        norm = [normalize_text(c) for c in row]
        if any("data" in c for c in norm) and any("valor" in c or "movimentacao" in c for c in norm):
            header_idx = idx
            col_indices = _map_flash_columns(norm)
            break

    # If no named header, detect by content
    if not col_indices:
        data_sample = all_rows[:10]
        col_indices = _detect_columns_by_content(data_sample)
        start_row = 0
    else:
        start_row = header_idx + 1

    transactions: list[FlashTransaction] = []
    for idx in range(start_row, len(all_rows)):
        item = _build_transaction_from_row(all_rows[idx], col_indices, idx)
        if item is not None:
            transactions.append(item)

    return FlashStatement(transactions=tuple(transactions), file_name=file_name)


def parse_flash_xlsx(content: bytes, file_name: str = "extrato.xlsx") -> FlashStatement:
    """Parse Excel workbook exported from the Flash web portal or manual sheets."""
    import warnings
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(io.BytesIO(content), data_only=True)

    # Pick the sheet with the most populated rows
    best_sheet = wb.active
    max_non_empty = 0
    for ws in wb.worksheets:
        cnt = sum(1 for r in ws.iter_rows(values_only=True) if any(r))
        if cnt > max_non_empty:
            max_non_empty = cnt
            best_sheet = ws

    rows = list(best_sheet.iter_rows(values_only=True))
    if not rows:
        return FlashStatement(transactions=(), file_name=file_name)

    # Find header row
    header_idx = -1
    col_indices: dict[str, int] = {}
    for idx, row in enumerate(rows[:15]):
        str_row = ["" if c is None else str(c).strip() for c in row]
        norm = [normalize_text(c) for c in str_row]
        if any("data" in c for c in norm) and any("valor" in c or "movimentacao" in c for c in norm):
            header_idx = idx
            col_indices = _map_flash_columns(norm)
            break

    if not col_indices:
        data_sample = [[str(c) if c is not None else "" for c in r] for r in rows[:10]]
        col_indices = _detect_columns_by_content(data_sample)
        start_row = 0
    else:
        start_row = header_idx + 1

    transactions: list[FlashTransaction] = []
    for idx in range(start_row, len(rows)):
        item = _build_transaction_from_row(rows[idx], col_indices, idx)
        if item is not None:
            transactions.append(item)

    return FlashStatement(transactions=tuple(transactions), file_name=file_name)


def parse_flash_statement(content: bytes | str, file_name: str = "extrato.csv") -> FlashStatement:
    """Auto-detect format (CSV or Excel) and parse the Flash statement."""
    lower = file_name.lower()
    if lower.endswith((".xlsx", ".xls")):
        if isinstance(content, str):
            content = content.encode("utf-8")
        return parse_flash_xlsx(content, file_name)
    return parse_flash_csv(content, file_name)


def merge_flash_statements(statements: Sequence[FlashStatement]) -> FlashStatement:
    """Combine transactions from multiple statements into a single unified statement."""
    non_empty = [s for s in statements if s and s.transactions]
    if not non_empty:
        return FlashStatement(transactions=(), file_name="")
    if len(non_empty) == 1:
        return non_empty[0]

    all_txs: list[FlashTransaction] = []
    seen: set[tuple[object, ...]] = set()
    names = [s.file_name for s in non_empty if s.file_name]

    for st in non_empty:
        for tx in st.transactions:
            # Deduplicate exact duplicate rows across multiple statements
            key = (
                tx.purchase_date,
                tx.raw_date,
                tx.amount,
                normalize_text(tx.description),
                normalize_text(tx.person),
                tx.payment_method,
            )
            if key in seen:
                continue
            seen.add(key)
            all_txs.append(tx)

    # Sort descending by purchase date
    all_txs.sort(key=lambda t: t.purchase_date, reverse=True)

    reindexed = tuple(
        FlashTransaction(
            row_index=idx + 1,
            purchase_date=t.purchase_date,
            raw_date=t.raw_date,
            description=t.description,
            amount=t.amount,
            person=t.person,
            payment_method=t.payment_method,
            statement_status=t.statement_status,
            is_deposit=t.is_deposit,
            hotel=t.hotel,
        )
        for idx, t in enumerate(all_txs)
    )

    combined_name = f"{len(non_empty)} arquivos (" + ", ".join(names[:3]) + ("..." if len(names) > 3 else "") + ")"
    return FlashStatement(transactions=reindexed, file_name=combined_name)

