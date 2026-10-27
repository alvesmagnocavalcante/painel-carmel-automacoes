"""Export services for Flash card reconciliation: Excel spreadsheet and email draft."""

from __future__ import annotations

import io
from collections import defaultdict
from decimal import Decimal
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from automations.cartao_flash.data import (
    STATUS_PENDING_LAUNCH,
    STATUS_PENDING_SETTLEMENT,
    STATUS_SETTLED,
    FlashReconciliationItem,
)


def export_flash_reconciliation(items: Sequence[FlashReconciliationItem]) -> bytes:
    """Export reconciliation items to Excel workbook following the company format."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência Flash"

    # Header styling
    header_fill = PatternFill(start_color="2B2724", end_color="2B2724", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    headers = [
        "MÊS",
        "DATA",
        "MOVIMENTAÇÃO",
        "VALOR",
        "PORTADOR",
        "OBSERVAÇÃO",
        "HOTEL",
        "DOC CMFLEX",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Status fonts and fills
    red_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    orange_font = Font(name="Calibri", size=10, bold=True, color="D97706")
    green_font = Font(name="Calibri", size=10, color="059669")
    normal_font = Font(name="Calibri", size=10)

    for row_idx, item in enumerate(items, start=2):
        tx = item.transaction
        date_str = tx.purchase_date.strftime("%d/%m/%Y")
        amount_val = -float(tx.amount)  # Displayed as negative in the team's spreadsheet
        cmflex_info = f"Doc #{item.matched_document.launch_number}" if item.matched_document else ""

        ws.append([
            tx.month_name,
            date_str,
            tx.description,
            amount_val,
            tx.person,
            item.effective_notes,
            item.effective_hotel,
            cmflex_info,
        ])

        # Style cells
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="left")

        val_cell = ws.cell(row=row_idx, column=4)
        val_cell.number_format = '"- R$ "#,##0.00;[Red]"- R$ "#,##0.00'
        val_cell.alignment = Alignment(horizontal="right")

        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="left")

        status_cell = ws.cell(row=row_idx, column=6)
        status_cell.alignment = Alignment(horizontal="center")
        if item.status == STATUS_PENDING_LAUNCH:
            status_cell.font = red_font
        elif item.status == STATUS_PENDING_SETTLEMENT:
            status_cell.font = orange_font
        else:
            status_cell.font = green_font

        ws.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=8).alignment = Alignment(horizontal="center")

        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = thin_border
            if c != 6:
                ws.cell(row=row_idx, column=c).font = normal_font

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_email_text(items: Sequence[FlashReconciliationItem]) -> str:
    """Generate email body text with pending items grouped by hotel and employee."""
    pendings = [item for item in items if item.status != STATUS_SETTLED]
    if not pendings:
        return (
            "Boa tarde,\n\n"
            "Conferência do cartão Flash realizada com sucesso: nenhuma pendência identificada!\n"
        )

    by_hotel: dict[str, dict[str, list[FlashReconciliationItem]]] = defaultdict(lambda: defaultdict(list))
    for item in pendings:
        hotel = item.effective_hotel
        person = item.transaction.person or "Portador não identificado"
        by_hotel[hotel][person].append(item)

    lines: list[str] = [
        "Boa tarde,",
        "",
        "Seguem as pendências de prestação do cartão Flash atualizadas para regularização:",
        "",
    ]

    for hotel, persons in sorted(by_hotel.items()):
        lines.append(f"🏨 HOTEL: {hotel.upper()}")
        lines.append("=" * 60)
        for person, user_items in sorted(persons.items()):
            lines.append(f"\n👤 Portador(a): {person}")
            for it in user_items:
                dt = it.transaction.purchase_date.strftime("%d/%m/%Y")
                val = f"R$ {it.transaction.amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                lines.append(f"  • {dt} | {it.transaction.description:<35} | {val:>10} | {it.effective_notes}")
        lines.append("\n")

    lines.append("Solicitamos aos responsáveis dos setores que prestem contas e realizem os lançamentos pendentes.")
    lines.append("\nAtenciosamente,\nContas a Pagar • Carmel Hotéis")
    return "\n".join(lines)
