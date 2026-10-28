"""Excel export and collection email generation for DDA reconciliation."""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from automations.dda.data import (
    STATUS_LANC,
    STATUS_LANC_OC,
    STATUS_NAO_LANC,
    DdaReconciliationItem,
)


def export_dda_reconciliation(items: Sequence[DdaReconciliationItem]) -> bytes:
    """Generate Excel workbook matching Carmel DDA reconciliation format (Cols A to O)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DDA Conferido"

    header_fill = PatternFill(start_color="2C241E", end_color="2C241E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    headers = [
        "EMPRESA",
        "BANCO",
        "CNPJ",
        "FORNECEDOR",
        "DOCUMENTO",
        "TIPO",
        "SITUAÇÃO",
        "DATA EMISSÃO",
        "DATA VENCIMENTO",
        "VALOR",
        "VALOR",
        "FORMULA",
        "STATUS",
        "OBSERVAÇÃO",
        "RESPONSAVEL",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Status fills and fonts
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name="Calibri", size=10, bold=True, color="15803D")

    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    yellow_font = Font(name="Calibri", size=10, bold=True, color="B45309")

    obs_highlight_fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    obs_highlight_font = Font(name="Calibri", size=10, bold=True, color="854D0E")

    normal_font = Font(name="Calibri", size=10)

    for row_idx, item in enumerate(items, start=2):
        tx = item.transaction
        doc = item.matched_document
        dt_emissao_str = tx.issue_date.strftime("%d/%m/%Y") if tx.issue_date else ""
        dt_venc_str = tx.due_date.strftime("%d/%m/%Y") if tx.due_date else ""
        amount_val = float(tx.amount)

        # Formula column emulation
        formula_val = amount_val if doc else "#N/D"

        status_val = item.effective_status
        obs_val = item.effective_observation
        resp_val = item.effective_responsible

        ws.append([
            tx.company,
            tx.bank_code,
            tx.formatted_cnpj,
            tx.supplier,
            tx.document_number,
            tx.charge_type,
            tx.situation,
            dt_emissao_str,
            dt_venc_str,
            amount_val,
            amount_val,
            formula_val,
            status_val,
            obs_val,
            resp_val,
        ])

        # Style row cells
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="left")
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="left")
        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=7).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=8).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=9).alignment = Alignment(horizontal="center")

        # Values format
        val_cell_1 = ws.cell(row=row_idx, column=10)
        val_cell_1.number_format = '"R$ "#,##0.00'
        val_cell_1.alignment = Alignment(horizontal="right")

        val_cell_2 = ws.cell(row=row_idx, column=11)
        val_cell_2.number_format = '"R$ "#,##0.00'
        val_cell_2.alignment = Alignment(horizontal="right")

        form_cell = ws.cell(row=row_idx, column=12)
        if isinstance(formula_val, (int, float)):
            form_cell.number_format = '"R$ "#,##0.00'
        form_cell.alignment = Alignment(horizontal="center")

        # Status cell styling
        status_cell = ws.cell(row=row_idx, column=13)
        status_cell.alignment = Alignment(horizontal="center")
        if status_val == STATUS_LANC:
            status_cell.fill = green_fill
            status_cell.font = green_font
        elif status_val == STATUS_NAO_LANC:
            status_cell.fill = red_fill
            status_cell.font = red_font
        elif status_val == STATUS_LANC_OC:
            status_cell.fill = yellow_fill
            status_cell.font = yellow_font
        else:
            status_cell.font = normal_font

        # Observation styling
        obs_cell = ws.cell(row=row_idx, column=14)
        obs_cell.alignment = Alignment(horizontal="left")
        if obs_val and obs_val not in ("SEM ALTERAÇÃO", "OK"):
            obs_cell.fill = obs_highlight_fill
            obs_cell.font = obs_highlight_font
        else:
            obs_cell.font = normal_font

        # Responsible styling
        resp_cell = ws.cell(row=row_idx, column=15)
        resp_cell.alignment = Alignment(horizontal="left")
        resp_cell.font = Font(name="Calibri", size=10, bold=bool(resp_val))

        for c_idx in range(1, 16):
            cell = ws.cell(row=row_idx, column=c_idx)
            cell.border = thin_border
            if not cell.font or cell.font == Font():
                cell.font = normal_font

    # Column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt_money(val: Decimal | float | int) -> str:
    """Format numeric value as Brazilian currency string R$ 1.234,56."""
    return f"R$ {float(val):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def generate_collection_email(items: Sequence[DdaReconciliationItem]) -> str:
    """Generate collection email text grouping pending and unlaunched items by responsible."""
    # Filter only items needing action
    action_items = [
        it for it in items
        if it.effective_status != STATUS_LANC or (
            it.effective_observation and it.effective_observation not in ("SEM ALTERAÇÃO", "OK")
        )
    ]

    if not action_items:
        return "Nenhuma pendência encontrada nos lançamentos do DDA conferidos."

    by_responsible: dict[str, list[DdaReconciliationItem]] = defaultdict(list)
    for it in action_items:
        resp = it.effective_responsible or "NÃO IDENTIFICADO / A DEFINIR"
        by_responsible[resp].append(it)

    lines = [
        "Prezada equipe do Contas a Pagar,",
        "",
        "Segue o resumo de boletos do DDA pendentes de lançamento ou com pendências cadastrais/anexos:",
        "",
    ]

    total_pending_val = sum(it.transaction.amount for it in action_items)

    for resp, resp_items in sorted(by_responsible.items()):
        subtotal = sum(it.transaction.amount for it in resp_items)
        lines.append("═══════════════════════════════════════════════════")
        lines.append(f"👤 RESPONSÁVEL: {resp.upper()} ({len(resp_items)} títulos - {_fmt_money(subtotal)})")
        lines.append("═══════════════════════════════════════════════════")

        for it in resp_items:
            tx = it.transaction
            doc = it.matched_document
            venc_str = tx.due_date.strftime("%d/%m/%Y") if tx.due_date else "Sem venc."
            doc_str = f" [Doc CMFlex #{doc.launch_number}]" if doc else " [NÃO LANÇADO NO ERP]"
            obs_str = f" -> {it.effective_observation}" if it.effective_observation else ""

            lines.append(
                f" • Venc: {venc_str} | {_fmt_money(tx.amount)} | {tx.supplier} ({tx.formatted_cnpj}) "
                f"| Hotel: {tx.company or tx.sheet_name}{doc_str}{obs_str}"
            )
        lines.append("")

    lines.append("───────────────────────────────────────────────────")
    lines.append(f"Total geral com pendências: {len(action_items)} boletos | {_fmt_money(total_pending_val)}")
    lines.append("Por favor, verificar e regularizar as pendências acima junto aos fornecedores e compradores.")

    return "\n".join(lines)
