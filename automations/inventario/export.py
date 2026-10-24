"""In-memory Excel export with the same matrix and totals as the panel."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from automations.inventario.data import InventoryCount, build_matrix, months_between, month_label

HEADER_FILL = PatternFill("solid", fgColor="4C3427")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _safe_text(value: str) -> str:
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def style_header_cells(cells, *, alignment: Alignment | None = None) -> None:
    for cell in cells:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        if alignment is not None:
            cell.alignment = alignment


def apply_column_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_inventory(rows: tuple[InventoryCount, ...], start: date, end: date, company: str) -> bytes:
    months = months_between(start, end)
    workbook = Workbook()
    try:
        sheet = workbook.active
        sheet.title = "Inventário mensal"
        sheet.cell(1, 1, "Conferência de inventário • Carmel Hotéis")
        sheet.cell(2, 1, f"Período: {start:%d/%m/%Y} a {end:%d/%m/%Y}")
        sheet.cell(3, 1, f"Empresa: {company}")
        sheet.cell(4, 1, "Artigos distintos por mês e almoxarifado; diferença em qualquer registro do período.")
        sheet.merge_cells(start_row=6, start_column=1, end_row=7, end_column=1)
        sheet.cell(6, 1, "Empresa / Almoxarifado")
        for index, month in enumerate(months):
            column = 2 + index * 3
            sheet.merge_cells(start_row=6, start_column=column, end_row=6, end_column=column + 2)
            sheet.cell(6, column, month_label(month))
            for offset, label in enumerate(("Contados", "Com diferença", "% diferença")):
                sheet.cell(7, column + offset, label)
        center = Alignment(horizontal="center", vertical="center")
        for cells in sheet.iter_rows(min_row=6, max_row=7, max_col=1 + len(months) * 3):
            style_header_cells(cells, alignment=center)
        for row_number, line in enumerate(build_matrix(rows, months), 8):
            sheet.cell(row_number, 1, _safe_text(line.label))
            for index, item in enumerate(line.values):
                if item is None:
                    continue
                column = 2 + index * 3
                sheet.cell(row_number, column, item.counted)
                sheet.cell(row_number, column + 1, item.different)
                percent = sheet.cell(row_number, column + 2, float(item.percentage) if item.percentage is not None else None)
                percent.number_format = "0.00%"
            if line.kind != "warehouse":
                for cell in sheet[row_number]:
                    cell.fill = PatternFill("solid", fgColor="DBD4CA" if line.kind == "company" else "9F9385")
                    cell.font = Font(bold=True, color="151311")
        sheet.freeze_panes = "B8"
        sheet.column_dimensions["A"].width = 52
        for column in range(2, 2 + len(months) * 3):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        return workbook_bytes(workbook)
    finally:
        workbook.close()
