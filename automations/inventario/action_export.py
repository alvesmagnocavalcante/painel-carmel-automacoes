"""Create new Excel/PDF downloads from saved action plans; never edit templates."""

from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from automations.inventario.action_plan import STATUSES, parse_date
from automations.inventario.export import (
    HEADER_FILL,
    _safe_text,
    apply_column_widths,
    style_header_cells,
    workbook_bytes,
)

LOGO = Path(__file__).resolve().parents[2] / "portal" / "assets" / "carmel-logo.png"
HEADERS = ("Hotel", "Data do inventário", "O quê?", "Como?", "Responsável", "Início",
           "Término", "Novo prazo", "Data realizada", "% concluído", "% planejado", "Status", "Observações")
COLUMN_WIDTHS = {
    "A": 24, "B": 19, "C": 65, "D": 60, "E": 25, "F": 16, "G": 16, "H": 16,
    "I": 18, "J": 17, "K": 17, "L": 20, "M": 55, "O": 22, "P": 18,
}


def export_actions_excel(actions) -> bytes:
    if not actions:
        raise ValueError("Não existem ações na seleção para exportar.")
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "Plano de Ação"
        book.calculation = CalcProperties(fullCalcOnLoad=True)
        sheet.merge_cells("A1:M1")
        sheet["A1"] = "CARMEL HOTÉIS | PLANO DE AÇÃO - INVENTÁRIO"
        sheet["A1"].font = Font(name="Calibri", size=19, bold=True, color="FFFFFF")
        sheet["A1"].fill = HEADER_FILL
        sheet.row_dimensions[1].height = 36
        sheet.merge_cells("A2:M2")
        sheet["A2"] = "Pontos de melhoria da API e preenchimento manual salvo no painel."
        sheet["O1"], sheet["P1"] = "Data de referência", "=TODAY()"
        sheet["P1"].number_format = "dd/mm/yyyy"
        last = 8 + len(actions)
        for index, status in enumerate(STATUSES, 1):
            sheet.cell(4, index, status)
            sheet.cell(5, index, f'=COUNTIF($L$9:$L${last},{sheet.cell(4,index).coordinate})')
            sheet.cell(4, index).font = Font(bold=True, color="4C3427")
            sheet.cell(5, index).font = Font(size=18, bold=True, color="4C3427")
        for index, label in enumerate(HEADERS, 1):
            sheet.cell(8, index, label)
        style_header_cells(sheet[8][:len(HEADERS)], alignment=Alignment(wrap_text=True, vertical="center"))
        sheet.row_dimensions[8].height = 30
        for row, action in enumerate(actions, 9):
            manual = action.manual.validate()
            try:
                inventory_date = parse_date(action.inventory_date) or action.inventory_date
            except ValueError:
                inventory_date = _safe_text(action.inventory_date)
            values = (action.hotel, inventory_date, action.improvement, manual.how, manual.responsible,
                      parse_date(manual.start), parse_date(manual.due), parse_date(manual.revised),
                      parse_date(manual.completed), manual.progress / 100, None, None, manual.notes)
            for column, value in enumerate(values, 1):
                if isinstance(value, str):
                    if len(value) > 32767:
                        raise ValueError("Uma ação ultrapassa o limite de texto de uma célula Excel.")
                    value = _safe_text(value)
                cell = sheet.cell(row, column, value)
                cell.font = Font(name="Calibri", size=11, color="151311")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor="EFEBE4" if row % 2 else "FFFFFF")
            sheet.cell(row, 11, f'=IF(OR(F{row}="",G{row}=""),"",IF(G{row}<=F{row},IF($P$1>=G{row},1,0),MAX(0,MIN(1,($P$1-F{row})/(G{row}-F{row})))))')
            deadline = f'IF(H{row}<>"",H{row},G{row})'
            sheet.cell(row, 12, f'=IF(OR(I{row}<>"",J{row}>=1),"Concluída",IF(AND({deadline}<>"",{deadline}<$P$1),"Atrasada",IF(H{row}<>"","Replanejada",IF(OR(F{row}<>"",J{row}>0),"Em andamento","Não iniciada"))))')
            for column in (2, 6, 7, 8, 9):
                sheet.cell(row, column).number_format = "dd/mm/yyyy"
            for column in (10, 11):
                sheet.cell(row, column).number_format = "0%"
            longest = max(len(action.improvement), len(manual.how), len(manual.notes))
            sheet.row_dimensions[row].height = min(409, max(54, (longest // 65 + 2) * 15))
        progress = DataValidation(type="decimal", operator="between", formula1=0, formula2=1)
        progress.error = "Preencha um percentual entre 0% e 100%."
        progress.showErrorMessage = True
        sheet.add_data_validation(progress)
        progress.add(f"J9:J{last}")
        for status, color in (("Concluída", "D8EBDD"), ("Atrasada", "F6DDDA"), ("Replanejada", "FFF0CA")):
            sheet.conditional_formatting.add(f"L9:L{last}", FormulaRule(formula=[f'$L9="{status}"'], fill=PatternFill("solid", fgColor=color)))
        apply_column_widths(sheet, COLUMN_WIDTHS)
        sheet.freeze_panes = "C9"
        sheet.auto_filter.ref = f"A8:M{last}"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:8"
        sheet.print_area = f"A1:M{last}"
        return workbook_bytes(book)
    finally:
        book.close()


def export_actions_pdf(actions) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    if not actions:
        raise ValueError("Não existem ações na seleção para exportar.")
    output = BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("PlanBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9,
                          leading=13, textColor=colors.HexColor("#151311"), alignment=TA_LEFT)
    heading = ParagraphStyle("PlanHeading", parent=styles["Heading2"], fontSize=12,
                             leading=16, textColor=colors.HexColor("#4C3427"), spaceBefore=14, keepWithNext=True)
    paragraph = lambda value: Paragraph(escape(str(value or "Não informado")).replace("\n", "<br/>"), body)
    story = [Paragraph("Plano de ação - Inventário", styles["Title"]),
             paragraph(f"Carmel Hotéis | Emitido em {date.today():%d/%m/%Y} | {len(actions)} registros"), Spacer(1, 5 * mm)]
    if LOGO.is_file():
        logo = Image(str(LOGO), width=40 * mm, height=17 * mm, kind="proportional")
        logo.hAlign = "LEFT"
        story.insert(0, logo)
    for index, action in enumerate(actions, 1):
        manual = action.manual.validate()
        story.append(Paragraph(escape(f"{index:02d} · {action.hotel} | Inventário: {action.inventory_date}"), heading))
        details = [("O quê?", action.improvement), ("Como?", manual.how), ("Responsável", manual.responsible),
                   ("Prazos", f"Início: {manual.start or '—'} | Término: {manual.due or '—'}\nNovo prazo: {manual.revised or '—'} | Realizado: {manual.completed or '—'}"),
                   ("Andamento", f"{manual.progress}% concluído | {manual.status()}"), ("Observações", manual.notes)]
        table = LongTable([[paragraph(label), paragraph(value)] for label, value in details],
                          colWidths=[30 * mm, 144 * mm], splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#DBD4CA")),
            ("ROWBACKGROUNDS", (1, 0), (-1, -1), [colors.HexColor("#F5F2ED"), colors.white]),
            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LINEBELOW", (0, 0), (-1, -1), .4, colors.HexColor("#DBD4CA")),
        ]))
        story.extend([table, Spacer(1, 3 * mm)])

    def page_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#4C3427"))
        canvas.drawString(18 * mm, 12 * mm, "Carmel Hotéis | Plano de ação de inventário")
        canvas.drawRightString(192 * mm, 12 * mm, f"Página {doc.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=22 * mm, title="Plano de ação - Inventário", author="Carmel Hotéis")
    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return output.getvalue()
