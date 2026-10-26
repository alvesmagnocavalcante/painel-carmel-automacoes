"""Planilha da conferência, sem alterar a fonte do CMFlex."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from automations.cartoes_credito.data import CardPayment
from automations.cartoes_credito.reconciliation import CARD_CHARGE_STATUS, BankReview
from automations.cartoes_credito.statement import BankStatement


HEADERS = (
    "Empresa", "Lançamento", "Favorecido", "CPF/CNPJ", "Valor", "Data do lançamento",
    "Vencimento", "Data programada", "Forma de pagamento", "Anexos",
    "Conferência", "Pendência", "Responsável", "Histórico",
)


def _safe_text(value: str) -> str:
    # A beneficiary or free-text history may begin with a spreadsheet formula.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_payments(rows: tuple[CardPayment, ...]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cartões de crédito"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4C3427")
    for row in rows:
        sheet.append((
            _safe_text(row.company), _safe_text(row.launch_number), _safe_text(row.beneficiary),
            _safe_text(row.tax_id), row.amount, row.launch_date, row.due_date, row.scheduled_date,
            _safe_text(row.payment_method), row.attachment_count, row.situation,
            _safe_text(row.pending_reason), _safe_text(row.responsible), _safe_text(row.history),
        ))
        index = sheet.max_row
        sheet.cell(index, 5).number_format = '"R$ "#,##0.00'
        for column in (6, 7, 8):
            sheet.cell(index, column).number_format = "DD/MM/YYYY"
        color = {"Conferido": "D9EAD3", "Pendente": "F4CCCC"}.get(row.situation)
        if color:
            for cell in sheet[index]:
                cell.fill = PatternFill("solid", fgColor=color)
    widths = (23, 17, 38, 21, 16, 19, 17, 19, 23, 10, 16, 42, 28, 54)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_statement(statement: BankStatement, reviews: dict[int, BankReview],
                     payments: tuple[CardPayment, ...]) -> bytes:
    """Create an auditable copy; never modify the imported bank workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Conferência da fatura"
    headers = ("Linha origem", "Titular", "Final do cartão", "Tipo", "Data na fatura",
               "Descrição na fatura", "Valor na fatura", "Empresa CMFlex", "Lançamento CMFlex",
               "Favorecido CMFlex", "Data lançamento CMFlex", "Valor CMFlex", "Anexos",
               "Classificação", "Pendência")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4C3427")
    payment_by_key = {payment.key: payment for payment in payments}
    for transaction in statement.transactions:
        review = reviews.get(transaction.row_number, BankReview())
        payment = payment_by_key.get(review.payment_key)
        sheet.append((
            transaction.row_number, _safe_text(transaction.cardholder), transaction.card_last_four,
            transaction.category, transaction.purchase_date, _safe_text(transaction.description),
            transaction.amount, _safe_text(payment.company) if payment else None,
            _safe_text(payment.launch_number) if payment else None,
            _safe_text(payment.beneficiary) if payment else None,
            payment.launch_date if payment else None, payment.amount if payment else None,
            payment.attachment_count if payment else None, review.status, _safe_text(review.reason),
        ))
        index = sheet.max_row
        for column in (5, 11):
            sheet.cell(index, column).number_format = "DD/MM/YYYY"
        for column in (7, 12):
            sheet.cell(index, column).number_format = '"R$ "#,##0.00;[Red]("R$ "#,##0.00)'
        color = {"Conferido": "D9EAD3", CARD_CHARGE_STATUS: "D9EAD3",
                 "Pendente": "F4CCCC"}.get(review.status)
        if color:
            for cell in sheet[index]:
                cell.fill = PatternFill("solid", fgColor=color)
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 6, "Total da fatura")
    sheet.cell(total_row, 7, f"=SUM(G2:G{total_row - 1})")
    sheet.cell(total_row, 7).number_format = '"R$ "#,##0.00;[Red]("R$ "#,##0.00)'
    sheet.cell(total_row + 1, 6, "Total informado pelo Itaú")
    sheet.cell(total_row + 1, 7, statement.total)
    sheet.cell(total_row + 1, 7).number_format = '"R$ "#,##0.00;[Red]("R$ "#,##0.00)'
    widths = (13, 29, 16, 17, 18, 39, 19, 25, 20, 36, 24, 18, 10, 17, 43)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:O{total_row - 1}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
