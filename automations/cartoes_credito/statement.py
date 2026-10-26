"""Leitura da planilha de fatura exportada pelo Itaú, sem guardar o arquivo."""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


MAX_STATEMENT_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_TRANSACTIONS = 1000


@dataclass(frozen=True)
class BankTransaction:
    row_number: int
    purchase_date: date
    description: str
    amount: Decimal
    cardholder: str
    card_last_four: str
    category: str


@dataclass(frozen=True)
class BankStatement:
    company: str
    due_date: date
    total: Decimal
    transactions: tuple[BankTransaction, ...]


def _normalize(value: object) -> str:
    text = str(value or "").casefold()
    return "".join(char for char in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(char))


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _amount(value: object) -> Decimal | None:
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    return None


def parse_rows(rows: list[tuple[object, ...]]) -> BankStatement:
    if len(rows) < 30 or not _normalize(rows[5][0]).startswith("itau empresas"):
        raise ValueError("A planilha não tem o formato de fatura Itaú esperado.")
    company = str(rows[2][0] or "").strip()
    due_date = _date(rows[8][2])
    total = _amount(rows[15][10])
    national_total = _amount(rows[12][10])
    international_total = _amount(rows[13][10])
    charges_total = _amount(rows[14][10])
    if not company or due_date is None or None in (total, national_total, international_total, charges_total):
        raise ValueError("Empresa, vencimento ou totais da fatura não foram encontrados.")

    holder = last_four = ""
    category = ""
    transactions: list[BankTransaction] = []
    totals = {"Nacional": Decimal("0"), "Internacional": Decimal("0"), "Encargo": Decimal("0")}
    for number, row in enumerate(rows, 1):
        first = row[0] if row else None
        normalized = _normalize(first)
        if isinstance(first, str):
            if " - final " in normalized:
                holder = first.split(" - FINAL", 1)[0].strip()
                match = re.search(r"(\d{4})\s*$", first)
                last_four = match.group(1) if match else ""
                category = ""
            elif normalized.startswith("lancamentos nacionais"):
                category = "Nacional" if holder else ""
            elif normalized.startswith("lancamentos internacionais"):
                category = "Internacional" if holder else ""
            elif normalized.startswith("produtos, servicos e encargos"):
                category = "Encargo" if holder else ""
            elif normalized.startswith("total de ") or normalized.startswith("encargos desta fatura"):
                category = ""
        value = _amount(row[10]) if len(row) > 10 else None
        purchase_date = _date(first)
        description = row[2] if len(row) > 2 else None
        if category and purchase_date and isinstance(description, str) and description.strip() and value is not None:
            transactions.append(BankTransaction(number, purchase_date, description.strip(), value,
                                                holder, last_four, category))
            totals[category] += value
            if len(transactions) > MAX_TRANSACTIONS:
                raise ValueError("A fatura excede o limite de 1.000 transações.")

    if not transactions:
        raise ValueError("Nenhuma transação da fatura foi encontrada.")
    expected = {"Nacional": national_total, "Internacional": international_total, "Encargo": charges_total}
    if any(abs(totals[key] - expected[key]) > Decimal("0.01") for key in expected):
        raise ValueError("Os lançamentos extraídos não conferem com os subtotais da fatura.")
    if abs(sum(totals.values()) - total) > Decimal("0.01"):
        raise ValueError("Os lançamentos extraídos não conferem com o total da fatura.")
    return BankStatement(company, due_date, total, tuple(transactions))


def parse_statement(content: bytes) -> BankStatement:
    if not content or len(content) > MAX_STATEMENT_BYTES:
        raise ValueError("Selecione uma planilha XLSX de até 5 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("A planilha descompactada excede 30 MB.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("A fatura deve ter uma única aba.")
            sheet = workbook.active
            if sheet.max_row > 5000 or sheet.max_column > 30:
                raise ValueError("A fatura excede o tamanho esperado.")
            rows = [tuple(row) for row in sheet.iter_rows(min_col=1, max_col=11, values_only=True)]
            return parse_rows(rows)
        finally:
            workbook.close()
    except (BadZipFile, OSError, KeyError) as error:
        raise ValueError("A planilha XLSX está inválida ou corrompida.") from error
